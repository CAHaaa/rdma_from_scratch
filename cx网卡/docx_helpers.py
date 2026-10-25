# -*- coding: utf-8 -*-
"""Reusable DOCX helpers for the ConnectX-6 RoCE guide."""
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

NAVY = "0B2545"
BLUE = "1F4E79"
TEAL = "2A9D8F"
ORANGE = "B45309"
RED = "B42318"
GRAY = "5B6B7C"
LIGHT_BLUE = "EEF4FA"
LIGHT_GRAY = "F4F6F8"
LIGHT_GREEN = "EEF8F5"
LIGHT_ORANGE = "FFF6E8"
LIGHT_RED = "FDECEC"
WHITE = "FFFFFF"
BORDER = "D6DEE7"
BODY_FONT = "DengXian"
HEAD_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"


def rgb(hex_str):
    return RGBColor.from_string(hex_str)


def set_run_font(run, name=BODY_FONT, size=10.5, bold=False, color=None, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name if name in ("Microsoft YaHei", "DengXian") else "Microsoft YaHei")
    rFonts.set(qn("w:cs"), name)
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rPr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")


def set_paragraph_format(p, before=0, after=4, line=1.35, align=None, keep_next=False,
                         first_indent=None, left_indent=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if keep_next:
        pf.keep_with_next = True
    if first_indent is not None:
        pf.first_line_indent = Cm(first_indent)
    if left_indent is not None:
        pf.left_indent = Cm(left_indent)


_TOKEN_RE = re.compile(r"(\*\*.*?\*\*|\[S\d+\])")


def add_rich_text(p, text, size=10.5, color=NAVY, bold=False, font=BODY_FONT):
    """Add text with **bold** and [S#] citation tokens."""
    for part in _TOKEN_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = p.add_run(part[2:-2])
            set_run_font(r, font, size, True, color)
        elif re.fullmatch(r"\[S\d+\]", part):
            r = p.add_run(part)
            set_run_font(r, font, max(8.0, size - 1.4), False, BLUE)
        else:
            r = p.add_run(part)
            set_run_font(r, font, size, bold, color)


def P(doc, text, size=10.5, color=NAVY, indent=True, before=0, after=5, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=before, after=after, line=1.38, align=align,
                         first_indent=0.74 if indent else None)
    add_rich_text(p, text, size, color)
    return p


def BULLETS(doc, items, size=10.2, color=NAVY, before=0, after=2):
    out = []
    for item in items:
        p = doc.add_paragraph()
        set_paragraph_format(p, before=before, after=after, line=1.30,
                             align=WD_ALIGN_PARAGRAPH.JUSTIFY, left_indent=0.55)
        p.paragraph_format.first_line_indent = Cm(-0.38)
        r = p.add_run("• ")
        set_run_font(r, BODY_FONT, size, True, BLUE)
        add_rich_text(p, item, size, color)
        out.append(p)
    return out


def NUMBERED(doc, items, size=10.2, color=NAVY, before=0, after=2):
    out = []
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        set_paragraph_format(p, before=before, after=after, line=1.30,
                             align=WD_ALIGN_PARAGRAPH.JUSTIFY, left_indent=0.65)
        p.paragraph_format.first_line_indent = Cm(-0.48)
        r = p.add_run(f"{i}. ")
        set_run_font(r, BODY_FONT, size, True, BLUE)
        add_rich_text(p, item, size, color)
        out.append(p)
    return out


def H1(doc, text, number=None):
    label = f"{number}  {text}" if number else text
    p = doc.add_heading(level=1)
    r = p.add_run(label)
    set_run_font(r, HEAD_FONT, 18, True, NAVY)
    set_paragraph_format(p, before=0, after=12, line=1.15, keep_next=True)
    p.paragraph_format.page_break_before = True
    # bottom rule
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), TEAL)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def H2(doc, text, number=None):
    label = f"{number}  {text}" if number else text
    p = doc.add_heading(level=2)
    r = p.add_run(label)
    set_run_font(r, HEAD_FONT, 13.5, True, BLUE)
    set_paragraph_format(p, before=14, after=6, line=1.2, keep_next=True)
    return p


def H3(doc, text):
    p = doc.add_heading(level=3)
    r = p.add_run(text)
    set_run_font(r, HEAD_FONT, 11.3, True, NAVY)
    set_paragraph_format(p, before=10, after=4, line=1.2, keep_next=True)
    return p


def CAPTION(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=2, after=8, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_rich_text(p, text, 8.5, GRAY)
    return p


def set_cell_background(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_margins(cell, top=0.06, bottom=0.06, left=0.10, right=0.10):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        node = OxmlElement(f"w:{k}")
        node.set(qn("w:w"), str(int(v * 567)))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcPr.append(mar)


def set_cell_borders(cell, color=BORDER, size=4, left=None):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        if left and edge == "left":
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(left))
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def set_table_borders(table, color=BORDER, size=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement("w:cantSplit")
    trPr.append(cantSplit)


def set_cell_width(cell, cm):
    cell.width = Cm(cm)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(cm * 567)))
    tcW.set(qn("w:type"), "dxa")


def TABLE(doc, headers, rows, widths, font_size=8.6, header_fill=NAVY,
          zebra=True, first_col_bold=False, align_center=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, BORDER, 4)
    # Explicit grid widths are required for Word/LibreOffice to honor column widths.
    for i, w in enumerate(widths):
        table.columns[i].width = Cm(w)
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_width(cell, widths[i])
        set_cell_background(cell, header_fill)
        set_cell_margins(cell, 0.08, 0.08, 0.10, 0.10)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        set_paragraph_format(p, before=0, after=0, line=1.12,
                             align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(h)
        set_run_font(r, HEAD_FONT, font_size + 0.2, True, WHITE)
    for ridx, row in enumerate(rows):
        tr = table.add_row()
        set_row_cant_split(tr)
        for i, val in enumerate(row):
            cell = tr.cells[i]
            set_cell_width(cell, widths[i])
            set_cell_margins(cell, 0.06, 0.06, 0.10, 0.10)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if zebra and ridx % 2 == 1:
                set_cell_background(cell, "F7F9FC")
            p = cell.paragraphs[0]
            set_paragraph_format(p, before=0, after=0, line=1.16,
                                 align=WD_ALIGN_PARAGRAPH.CENTER if align_center else WD_ALIGN_PARAGRAPH.LEFT)
            txt = str(val)
            # allow explicit line breaks
            if "\n" in txt:
                parts = txt.split("\n")
                for j, part in enumerate(parts):
                    if j:
                        p.add_run().add_break()
                    if i == 0 and first_col_bold:
                        r = p.add_run(part)
                        set_run_font(r, BODY_FONT, font_size, True, NAVY)
                    else:
                        add_rich_text(p, part, font_size, NAVY)
            else:
                if i == 0 and first_col_bold:
                    r = p.add_run(txt)
                    set_run_font(r, BODY_FONT, font_size, True, NAVY)
                else:
                    add_rich_text(p, txt, font_size, NAVY)
    return table


def CODE(doc, code, size=8.0, fill=LIGHT_GRAY):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, "E2E8F0", 4)
    set_row_cant_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_width(cell, 16.0)
    set_cell_background(cell, fill)
    set_cell_margins(cell, 0.10, 0.10, 0.14, 0.14)
    set_cell_borders(cell, "D6DEE7", 4, left=18)
    p = cell.paragraphs[0]
    set_paragraph_format(p, before=0, after=0, line=1.04)
    lines = code.strip("\n").split("\n")
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, CODE_FONT, size, False, "1F2937")
    return table


def CALLOUT(doc, title, body, kind="info"):
    palette = {
        "info": (LIGHT_BLUE, BLUE),
        "warn": (LIGHT_ORANGE, ORANGE),
        "risk": (LIGHT_RED, RED),
        "good": (LIGHT_GREEN, TEAL),
        "gray": (LIGHT_GRAY, GRAY),
    }
    fill, accent = palette.get(kind, palette["info"])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(2)
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = 1.0
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_row_cant_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_width(cell, 16.0)
    set_cell_background(cell, fill)
    set_cell_margins(cell, 0.12, 0.12, 0.18, 0.16)
    set_cell_borders(cell, accent, 4, left=24)
    p = cell.paragraphs[0]
    set_paragraph_format(p, before=0, after=2, line=1.15)
    r = p.add_run(title)
    set_run_font(r, HEAD_FONT, 9.8, True, accent)
    p2 = cell.add_paragraph()
    set_paragraph_format(p2, before=0, after=0, line=1.28,
                         align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_rich_text(p2, body, 9.3, NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def FIGURE(doc, path, width_cm=16.0):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=4, after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.add_run().add_picture(path, width=Cm(width_cm))


def add_hyperlink(paragraph, url, text, color=BLUE, underline=True, size=8.0):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), BODY_FONT)
    rFonts.set(qn("w:hAnsi"), BODY_FONT)
    rFonts.set(qn("w:eastAsia"), HEAD_FONT)
    rPr.append(rFonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rPr.append(c)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_field(run, field):
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = field
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "1"
    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(t)
    run._r.append(fldChar3)


def add_header_footer(section, title):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run(title)
    set_run_font(r, BODY_FONT, 8.0, False, GRAY)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run("第 ")
    set_run_font(r1, BODY_FONT, 8.0, False, GRAY)
    r2 = fp.add_run()
    set_run_font(r2, BODY_FONT, 8.0, False, GRAY)
    add_field(r2, "PAGE")
    r3 = fp.add_run(" 页")
    set_run_font(r3, BODY_FONT, 8.0, False, GRAY)


def start_page_numbering(section, start=1):
    sectPr = section._sectPr
    pgNumType = OxmlElement("w:pgNumType")
    pgNumType.set(qn("w:start"), str(start))
    sectPr.append(pgNumType)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(NAVY)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), BODY_FONT)
    rFonts.set(qn("w:hAnsi"), BODY_FONT)
    rFonts.set(qn("w:eastAsia"), BODY_FONT)
    rFonts.set(qn("w:cs"), BODY_FONT)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")
    rPr.append(lang)
    pf = normal.paragraph_format
    pf.line_spacing = 1.38
    pf.space_after = Pt(5)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        st = styles[name]
        st.font.name = HEAD_FONT
        st.font.color.rgb = rgb(NAVY)
        st.font.bold = True
        rPr = st.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:ascii"), HEAD_FONT)
        rFonts.set(qn("w:hAnsi"), HEAD_FONT)
        rFonts.set(qn("w:eastAsia"), HEAD_FONT)
