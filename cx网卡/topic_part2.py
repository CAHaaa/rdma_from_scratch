# -*- coding: utf-8 -*-
"""Chapters 6-11 and appendices for the CC/tuning companion."""
import os
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docx_helpers import (
    NAVY, BLUE, TEAL, ORANGE, RED, GRAY, WHITE, BORDER,
    LIGHT_BLUE, LIGHT_GRAY, LIGHT_GREEN, LIGHT_ORANGE, LIGHT_RED,
    BODY_FONT, HEAD_FONT,
    P, BULLETS, NUMBERED, H1, H2, H3, CAPTION, TABLE, CODE, CALLOUT, FIGURE,
    set_run_font, set_paragraph_format, add_rich_text, add_hyperlink,
    set_cell_background, set_cell_margins, set_cell_width,
    set_table_borders, set_repeat_table_header, set_row_cant_split,
    add_header_footer, start_page_numbering, setup_styles,
)
from common_topic import TOPIC_TITLE, TOPIC_SUBTITLE, TOPIC_VERSION, TOPIC_DATE, TOPIC_DOCX, TOPIC_SOURCES
from topic_part1 import add_cover, add_front_matter, add_section_1, add_section_2, add_section_3, add_section_4, add_section_5


def add_body_start(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.9)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)
    add_header_footer(sec, TOPIC_TITLE)
    start_page_numbering(sec, 1)


def add_section_6(doc):
    H1(doc, "其他算法：QCN、ZTR-RTT、PCC", "6")
    P(doc, "这一章把“与 CX6 相关但非主线”的算法讲清楚，目的不是让你都去用，而是避免把不同代际、不同平台的算法混在一起调。")

    H2(doc, "QCN（802.1Qau）：L2 的遗留方案", "6.1")
    P(doc, "QCN 是 IEEE 802.1Qau 定义的 Layer 2 拥塞通知：交换机检测到拥塞后生成 CNM（Congestion Notification Message），端点收到后降速。ConnectX-6 手册列有 802.1Qau 能力，但这不等于 RoCEv2 的 DCQCN。[C1][C4]")
    rows = [
        ("层", "Layer 2 / 不可路由", "DCQCN 在 IP/ECN 上，可跨 L3"),
        ("反馈", "CNM", "CNP"),
        ("主要场景", "Legacy L2 / RoCEv1", "RoCEv2 / 现代数据中心"),
        ("与 CX6 关系", "硬件能力存在", "RoCEv2 主线是 DCQCN"),
    ]
    TABLE(doc, ["维度", "QCN", "DCQCN"], rows, [3.0, 6.4, 6.6], font_size=8.4, first_col_bold=True)
    CALLOUT(doc, "不要把 QCN 当 DCQCN 调", "QCN 的 CNM 和 DCQCN 的 CNP 不是同一个反馈；QCN 的参数、计数器、交换机要求也不同。只有在既有 L2/QCN 环境中才考虑 QCN，RoCEv2 优先 DCQCN。", "warn")

    H2(doc, "ZTR-RTT：延迟型的新一代原生算法", "6.2")
    P(doc, "NVIDIA DOCA 3.5 把 ZTR-RTT 描述为 Zero Touch RoCE 的硬件拥塞控制算法：它同时使用 RTT 和 CNP，比较实测 RTT 与目标 RTT，按 AIMD 调整 cwnd/速率；通过硬件时间戳测量 RTT，减少对交换机 ECN 阈值配置的依赖。[C7]")
    BULLETS(doc, [
        "**信号**：RTT + CNP；RTT 由发送端/接收端硬件时间戳协作测量。[C7]",
        "**动作**：Additive Increase / Multiplicative Decrease；以 cwnd（字节）收敛到速率。[C7][C8]",
        "**特点**：首次拥塞快速反应；可减少 HOL 阻塞；目标是“zero touch”部署。[C7]",
        "**平台**：DOCA 将其放在 DPA-based CC engine/新一代原生 CC 基础设施中；不要假设 CX6 支持，按型号/固件文档确认。[C6][C7]",
    ])
    rows = [
        ("ALPHA", "目标 RTT 与 base RTT 的线性系数", "fxp16", "0.1"),
        ("MAX_DEC", "单次降速因子", "fxp16", "0.97"),
        ("MAX_INC", "单次升速因子", "fxp16", "1.06"),
        ("AI", "每 100Gb/s 的加性增量（字节）", "integer", "按平台"),
    ]
    TABLE(doc, ["参数", "含义", "类型", "默认/备注"], rows, [3.2, 7.6, 2.4, 2.8], font_size=8.2, first_col_bold=True)
    P(doc, "这些参数只用于说明 ZTR-RTT 的调节维度，**不是 CX6 DCQCN 的参数**。把 ALPHA/MAX_DEC/MAX_INC 套到 CX6 上没有意义。[C8]")

    H2(doc, "DOCA PCC：把算法变成可编程", "6.3")
    P(doc, "DOCA PCC 提供 RP/NP 的可编程接口：用户用 DPACC 编译自定义 CC 逻辑，在 BlueField-3 的 DPA 上执行；可配置 probe 包、CC 事件、每流速率限制等。[C9]")
    BULLETS(doc, [
        "**平台**：DOCA PCC 文档要求 BlueField-3（或更新）平台；不是 CX6 的即用功能。[C9]",
        "**启用**：RP 侧 `USER_PROGRAMMABLE_CC=1`，NP 侧 `PCC_INT_EN=0`；需要关机/重启。[C9]",
        "**开发**：需要 DPACC 编译、host library、device library；不是改几个 sysfs 就能完成。[C9]",
        "**适用**：研究/差异化算法、需要自定义信号（CCMAD/IFA）的场景；大多数 CX6 RoCE 部署用 DCQCN 更实际。[C9]",
    ])

    H2(doc, "TIMELY：RTT 型方案的代表", "6.4")
    P(doc, "TIMELY 用 RTT 作为拥塞信号，按 RTT 变化调整速率；它证明了“延迟信号”在数据中心可行，但不是 CX6 的原生配置项，也没有对应的 NVIDIA `mlxconfig` 参数。[C20]")
    P(doc, "ZTR-RTT 可以看作 NVIDIA 在新平台上把 RTT 信号工程化的一种实现；两者思想相近，但实现、参数、平台完全不同。不要把 TIMELY 的公式直接套到 DCQCN。")

    H2(doc, "一张表选算法", "6.5")
    rows = [
        ("DCQCN", "ECN + CNP", "CX6 / RoCEv2", "成熟、生态完整", "需 ECN/CNP QoS；PFC 可能掩盖拥塞", "默认选它"),
        ("QCN", "L2 CNM", "Legacy L2/RoCEv1", "L2 内闭环", "不可路由；与 DCQCN 不同", "仅遗留环境"),
        ("ZTR-RTT", "RTT + CNP", "新平台/DPA CC", "少依赖交换机 ECN 配置", "平台限制；参数体系不同", "按平台确认"),
        ("PCC", "自定义", "BlueField-3+", "最灵活", "开发/验证成本高", "研究/差异化"),
        ("TIMELY", "RTT", "研究/业界", "延迟信号", "非 CX6 原生", "理解概念"),
    ]
    TABLE(doc, ["算法", "信号", "平台", "优点", "边界", "建议"], rows, [2.4, 2.6, 3.2, 3.2, 3.0, 1.6], font_size=7.8, first_col_bold=True)


def add_section_7(doc):
    H1(doc, "参数敏感性与交互", "7")
    P(doc, "调参前先建立“方向感”：这个参数调大是让系统更保守还是更激进？它和谁耦合？如果不知道方向，实验就变成碰运气。")

    H2(doc, "敏感性矩阵", "7.1")
    rows = [
        ("ECN Kmin", "标记更晚；缓冲/时延更高；PFC 风险上升", "标记更早；吞吐可能下降；降速更频繁", "PFC 先触发/时延高", "ECN 标记率、队列时延"),
        ("ECN Kmax", "标记斜坡更长；同一队列下标记更弱", "标记更早变强；更激进", "过激降速", "Pmax、吞吐、CNP"),
        ("Pmax", "高队列时标记更强", "高队列时标记更弱；可能降速不足", "PFC/丢包", "标记率、队列"),
        ("MIN_TIME_BETWEEN_CNPS", "CNP 更少；反应更慢", "CNP 更密；控制面开销更高", "过冲/振荡", "CNP 速率、RP 降速"),
        ("DCE_TCP_G", "改变 α 平滑/敏感度；方向依赖实现", "同左；默认优先", "α 抖动/反应过慢", "α、CNP 到达模式"),
        ("RPG_GD", "单次降速更温和", "单次降速更狠", "降速不足/过激", "CNP 后速率曲线"),
        ("RPG_AI_RATE / HAI_RATE", "恢复更快；更易振荡", "恢复更慢；吞吐恢复差", "锯齿/恢复慢", "恢复时间、P99"),
        ("RPG_TIME_RESET / BYTE_RESET", "恢复更慢", "恢复更快", "恢复慢/再次拥塞", "恢复斜率"),
        ("PFC Xoff", "PFC 更晚；占用更多缓冲；headroom 风险", "PFC 更早；pause 更频繁", "丢包/风暴", "pause、buffer、drop"),
        ("Headroom / cable length", "吸收更多在途；占用更多缓冲", "缓冲更省；丢包风险上升", "PFC 前丢包", "no-buffer-discard"),
        ("Buffer pool", "吸收更强；时延更高", "更省；丢包风险", "时延/丢包", "buffer max usage"),
    ]
    TABLE(doc, ["参数", "调大", "调小", "主要风险", "先看"], rows, [3.4, 4.0, 4.0, 2.6, 2.0], font_size=7.8, first_col_bold=True)

    H2(doc, "四组关键耦合", "7.2")
    rows = [
        ("ECN Kmin ↔ PFC Xoff", "不同缓冲：ECN 在 Egress 队列，PFC 在 Ingress PG；不能直接比大小 [C12]", "目标是 ECN 先动作、PFC 兜底；用标记率/pause/丢包验证", "PFC 高 + ECN 低 → 先降 Kmin，不要先抬 Xoff"),
        ("CNP 速率 ↔ AI/HAI", "CNP 越密、恢复越猛，越容易振荡", "CNP 最小间隔、AI/HAI 一次只动一个", "降速后锯齿/过冲"),
        ("RPG_GD ↔ α/g", "两者都改变闭环增益", "不要同时调；先默认，再只动 RPG_GD", "降速过激/不足"),
        ("Headroom ↔ Xoff/cable/MTU", "Xoff 越高，留给在途报文的 headroom 越紧张", "按最坏 MTU/线缆/速率算 headroom", "PFC 配了仍丢包"),
    ]
    TABLE(doc, ["耦合", "为什么", "怎么处理", "典型现象"], rows, [3.2, 5.6, 4.6, 2.6], font_size=8.0, first_col_bold=True)

    H2(doc, "调优的“支配顺序”", "7.3")
    NUMBERED(doc, [
        "**正确性**：priority、trust、PFC 域、CNP strict、ECN TC——错一项，后面都白调。",
        "**缓冲/headroom**：先保证 PFC 不会在生效前丢包；否则测到的是缓冲问题，不是 CC 问题。",
        "**ECN 阈值**：Kmin/Kmax/Pmax 决定“何时给信号”，优先于端点参数。",
        "**CNP QoS/速率**：反馈通道是否可靠、是否过密。",
        "**DCQCN 恢复参数**：RPG_GD、AI/HAI、定时器——最后才动。",
        "**PFC 微调**：Xoff/Xon 是最后 5%，不要用它掩盖 ECN 问题。",
    ])
    CALLOUT(doc, "不要先动 RPG_GD", "RPG_GD 很显眼，但它只改变“单次降速幅度”。如果 ECN 没标记、CNP 没回来，调 RPG_GD 不会让闭环工作。先确认信号链，再调增益。", "warn")

    H2(doc, "不要这样调", "7.4")
    BULLETS(doc, [
        "不要同时改 Kmin、Pmax、RPG_GD、AI/HAI；无法归因。",
        "不要用“提高 Pmax”来解决“ECN 不触发”；Pmax 只影响高队列时的强度，不改变起始点。",
        "不要把 Xoff 调高来“减少 PFC”；Xoff 越高，PFC 越晚，越容易在生效前丢包。",
        "不要把 buffer 加到“永不丢包”；大缓冲会抬高时延、拖慢 ECN。",
        "不要用平均吞吐验收；P99、PFC pause、Jain 公平性同样重要。",
        "不要把 ZTR-RTT/PCC 的参数套到 CX6 DCQCN；两套体系不通用。",
    ])


def add_section_8(doc):
    H1(doc, "调优方法论：基线、实验与统计", "8")
    P(doc, "调优不是“改一个参数看吞吐”。它是一次受控实验：先定义目标，再建立基线，再改变一个因素，最后用分布和护栏指标判断是否真的更好。")

    H2(doc, "定义目标与护栏", "8.1")
    rows = [
        ("吞吐", "平均/总算力利用率、goodput", "达到线速的约定比例；incast 下可恢复", "不能以牺牲 P99/丢包为代价"),
        ("时延", "P50/P99/P999、队列时延", "P99 在 SLO 内", "平均值会掩盖长尾"),
        ("丢包", "no-buffer-discard、重传、out-of-sequence", "lossless priority 为 0 或接近 0", "任何持续丢包都是红线"),
        ("公平", "Jain、每流吞吐分布", "J 接近 1；无流饿死", "总吞吐正常也可能不公平"),
        ("稳定", "恢复时间、振荡幅度", "阶跃/incast 后能收敛", "只看稳态会漏掉振荡"),
        ("可运维", "PFC storm、配置持久化、告警", "无 storm；重启后自动恢复", "不能靠人工记忆"),
    ]
    TABLE(doc, ["维度", "指标", "目标示例", "护栏"], rows, [3.0, 6.2, 4.2, 2.6], font_size=8.2, first_col_bold=True)

    H2(doc, "基线", "8.2")
    BULLETS(doc, [
        "记录固件/驱动/交换机 NOS、PSID、端口速率、MTU、线缆长度、PFC/ECN/priority 配置。",
        "记录默认参数：`mlxconfig query`、`mlnx_qos -i`、交换机 `show roce`/`nv show qos roce`。[C16]",
        "记录业务基线：单流、4:1/8:1 incast、混合流量、故障切换。",
        "基线至少 3 次运行；每次有 warmup 和稳定态；不要用一次峰值当基线。",
    ])

    H2(doc, "负载矩阵", "8.3")
    rows = [
        ("单流", "1 发送端 → 1 接收端", "验证链路/MTU/GID/基本吞吐", "不能证明 CC 工作"),
        ("Incast 4:1", "4 发送端 → 1 接收端", "触发 ECN/CNP/PFC", "最常用的 CC 验收"),
        ("Incast 8:1/16:1", "更多发送端", "压力/公平性/恢复", "观察 PFC 风暴风险"),
        ("Permutation", "全互通/全对全", "公平性与整体吞吐", "接近 AI 训练 all-to-all"),
        ("Burst", "短时线速突发", "headroom/PFC", "平均速率低也可能触发"),
        ("Mixed", "RoCE + TCP/存储", "优先级隔离", "验证 PFC 不误伤"),
        ("Failure", "链路/端口/交换机故障", "收敛/回滚", "验证冗余与恢复"),
    ]
    TABLE(doc, ["场景", "拓扑/负载", "验证目标", "注意"], rows, [3.0, 5.2, 5.2, 2.6], font_size=8.0, first_col_bold=True)

    H2(doc, "实验设计", "8.4")
    NUMBERED(doc, [
        "**一次一个因素**：Kmin、Pmax、RPG_GD、AI/HAI 分开测；如果怀疑交互，用 2×2 矩阵，不要一次改四个。",
        "**设置水平**：默认值、±25%、±50%（或按参数范围取安全步长）；不要跨数量级跳。",
        "**重复**：每个水平至少 3 次；取中位数，记录 IQR/P99。",
        "**稳态**：warmup 后至少 5–10 分钟稳定态；incast 要覆盖同步突发周期。",
        "**盲测**：如果可能，让采集/分析人员不知道当前参数，减少确认偏差。",
        "**回滚**：每组实验前保存原值；出现护栏指标恶化立即回滚。",
    ])

    H2(doc, "统计与判断", "8.5")
    P(doc, "不要用“这次比上次高 2%”做结论。看效应量、分布和护栏：吞吐提升但 P99 翻倍，不是成功；P99 改善但丢包出现，也不是成功。")
    rows = [
        ("中心趋势", "中位数/均值", "至少 3 次", "均值受长尾影响"),
        ("尾部分布", "P99/P999", "每次运行", "AI/存储更关注长尾"),
        ("离散度", "IQR/标准差", "每次运行", "方差大说明不稳定"),
        ("公平性", "Jain 指数", "多流", "总吞吐正常也可能不公平"),
        ("护栏", "丢包/PFC storm/重传", "每次运行", "任何红线直接判失败"),
    ]
    TABLE(doc, ["维度", "统计量", "采样", "说明"], rows, [3.2, 4.4, 3.4, 5.0], font_size=8.2, first_col_bold=True)

    H2(doc, "决策规则", "8.6")
    BULLETS(doc, [
        "**保留**：吞吐/P99 改善达到预设效应量，且丢包/PFC/公平性没有恶化。",
        "**观察**：指标互有胜负；扩大样本或换负载再测。",
        "**回滚**：出现 lossless 丢包、PFC storm、P99 超 SLO、CNP 异常。",
        "**固化**：连续 3 次运行 + 重启测试通过，才写入启动配置。",
    ])
    CALLOUT(doc, "调优的停止条件", "出现持续 `no-buffer-discard`、PFC storm、CNP 计数异常下降、业务 SLO 明显恶化时，停止调参并回滚；不要“再试一个参数”。", "risk")


def add_section_9(doc):
    H1(doc, "分场景调优配方", "9")
    P(doc, "下面的配方是“起始假设”，不是通用答案。每个场景都给出：先看什么、先调什么、避免什么、用什么验收。")

    H2(doc, "AI 训练 / All-Reduce", "9.1")
    rows = [
        ("典型症状", "同步突发、incast、P99 尖刺；PFC pause 高；ECN 标记率随步同步波动"),
        ("先看", "ECN 标记率、PFC pause、CNP 计数、buffer max usage、GPU/NIC 利用率"),
        ("起点", "NVIDIA 默认 ECN/PFC/DCQCN；CNP strict；PFC 只开 RoCE priority"),
        ("优先调", "Kmin（按目标队列时延）→ CNP QoS → AI/HAI（仅在恢复慢时）"),
        ("避免", "用 PFC 硬扛 incast；把 Kmin 设得比 PFC 还晚；同时改 AI/HAI 和 RPG_GD"),
        ("验收", "4:1/8:1 incast、all-reduce 周期、P99、无 no-buffer-discard、无 storm"),
    ]
    TABLE(doc, ["项", "内容"], rows, [3.2, 12.8], font_size=8.2, first_col_bold=True)

    H2(doc, "存储 / 多 QP", "9.2")
    rows = [
        ("典型症状", "多 QP 并发、时延敏感；PFC 偶发；重传/乱序影响尾延迟"),
        ("先看", "PFC pause、ECN 标记、CNP、`out_of_sequence`、P99"),
        ("起点", "lossless 或 lossy 按业务；ECN 默认；CNP strict；PFC 只开 RoCE priority"),
        ("优先调", "ECN Kmin（按 RTT/队列时延）→ PFC headroom → DCQCN 恢复"),
        ("避免", "把存储和普通 TCP 混在一个 lossless priority；忽略慢接收端"),
        ("验收", "多 QP 并发、读写混合、P99、零丢包/低重传"),
    ]
    TABLE(doc, ["项", "内容"], rows, [3.2, 12.8], font_size=8.2, first_col_bold=True)

    H2(doc, "多租户 / Lossy", "9.3")
    rows = [
        ("典型症状", "跨 L3、租户共享、无法保证 PFC 端到端；允许重传"),
        ("先看", "ECN 标记、CNP、重传、租户间公平性"),
        ("起点", "ECN/DCQCN；不使用 PFC；CNP strict；按租户/优先级隔离"),
        ("优先调", "Kmin/Kmax 让 ECN 早于丢包；CNP QoS；RPG_GD 稳定降速"),
        ("避免", "跨 WAN/跨租户依赖 PFC；把 PFC 当 SLA 保证"),
        ("验收", "多租户并发、丢包可恢复、无租户饿死、P99 可接受"),
    ]
    TABLE(doc, ["项", "内容"], rows, [3.2, 12.8], font_size=8.2, first_col_bold=True)

    H2(doc, "混合 TCP + RoCE", "9.4")
    rows = [
        ("典型症状", "TCP 与 RoCE 共享 buffer/priority；PFC 误伤 TCP；RoCE 被 TCP 挤占"),
        ("先看", "priority 映射、PFC pause、ECN 标记、两类流量吞吐"),
        ("起点", "RoCE 独立 lossless priority；TCP 走 lossy priority；CNP strict"),
        ("优先调", "优先级/TC 隔离 → ECN → PFC 域"),
        ("避免", "所有 priority 都开 PFC；让 TCP 进入 lossless pool"),
        ("验收", "混合流量下两类业务均达标；PFC 不扩散到 TCP"),
    ]
    TABLE(doc, ["项", "内容"], rows, [3.2, 12.8], font_size=8.2, first_col_bold=True)

    H2(doc, "长 RTT / 多级 Fabric", "9.5")
    rows = [
        ("典型症状", "BDP 大、CNP 回来慢、阈值按短距配置导致过晚/过早"),
        ("先看", "RTT、BDP、队列时延、CNP 往返、PFC 域"),
        ("起点", "按 BDP/目标队列时延重算 Kmin；CNP strict；PFC 只到边缘"),
        ("优先调", "Kmin（时间单位）→ CNP 优先级 → 恢复参数"),
        ("避免", "直接复制 100 m/单跳的阈值；把 PFC 扩到跨 spine"),
        ("验收", "长 RTT incast、CNP 延迟、P99、无 PFC 风暴"),
    ]
    TABLE(doc, ["项", "内容"], rows, [3.2, 12.8], font_size=8.2, first_col_bold=True)

    H2(doc, "短流 / RPC", "9.6")
    rows = [
        ("典型症状", "流短、对首包/尾延迟敏感；ECN 可能来不及反应"),
        ("先看", "P99、PFC pause、ECN 标记、CNP 延迟"),
        ("起点", "ECN 早期标记；避免深缓冲；CNP strict；PFC 兜底"),
        ("优先调", "Kmin（较低）→ 缓冲/队列 → 恢复参数"),
        ("避免", "用大缓冲换吞吐；让短流排在长流后面"),
        ("验收", "短流 P99、混合长流、无饥饿"),
    ]
    TABLE(doc, ["项", "内容"], rows, [3.2, 12.8], font_size=8.2, first_col_bold=True)


def add_section_10(doc):
    H1(doc, "调优案例与决策树", "10")
    P(doc, "以下案例是“观察 → 假设 → 动作 → 验证”的模板。真实环境要先看计数器，不要因为症状相似就直接套动作。")

    H2(doc, "案例 1：ECN 标记高、吞吐低", "10.1")
    rows = [
        ("观察", "ECN marked 比例 > 80%，PFC pause 不高，吞吐低于 SLO"),
        ("假设", "Kmin 太低/Pmax 太高，端点降速过激；或拥塞确实严重"),
        ("动作", "先抬高 Kmin（让标记更晚），必要时降低 Pmax；一次只动一个"),
        ("验证", "标记率下降、吞吐恢复、P99 不恶化、无丢包"),
        ("如果无效", "看队列是否持续高：可能是容量/超配问题，不是 CC 参数"),
    ]
    TABLE(doc, ["项", "内容"], rows, [3.0, 13.0], font_size=8.3, first_col_bold=True)

    H2(doc, "案例 2：PFC pause 高、ECN 标记低", "10.2")
    rows = [
        ("观察", "prio3_tx_pause 持续高；ECN marked 很低；可能无丢包"),
        ("假设", "PFC 在替 ECN 干活：Kmin 太高或 ECN 未在正确 TC"),
        ("动作", "降低 Kmin（让 ECN 更早标记）；确认 ECN 在 RoCE TC；不要先抬高 Xoff"),
        ("验证", "ECN 标记上升、PFC pause 下降、吞吐/P99 稳定"),
        ("如果无效", "检查 ECN 是否真的启用、CNP 是否回得来、PFC 是否误配 priority"),
    ]
    TABLE(doc, ["项", "内容"], rows, [3.0, 13.0], font_size=8.3, first_col_bold=True)

    H2(doc, "案例 3：配了 PFC 仍然丢包", "10.3")
    rows = [
        ("观察", "`no-buffer-discard` 增长；PFC pause 有，但仍在 Xoff 后丢包"),
        ("假设", "headroom 不足、cable length 错、MTU/速率不匹配、Xoff 太高"),
        ("动作", "按最坏情况重算 headroom；核对 cable length/MTU；不要只加 Xoff"),
        ("验证", "no-buffer-discard 归零；PFC pause 可控；吞吐恢复"),
        ("如果无效", "查线缆/光模块/FEC；查是否同 priority 混入普通流量"),
    ]
    TABLE(doc, ["项", "内容"], rows, [3.0, 13.0], font_size=8.3, first_col_bold=True)

    H2(doc, "案例 4：ECN 标记了但 CNP 不回来", "10.4")
    rows = [
        ("观察", "交换机 ECN marked 增长；NP/CNP 计数不增长或很低；RP 不降速"),
        ("假设", "NP 未 enable、CNP 被丢/被 pause、DSCP/PCP 映射错、CNP 队列不 strict"),
        ("动作", "查 NP enable、CNP DSCP/PCP、TC6 strict、PFC 是否误开在 CNP priority"),
        ("验证", "CNP 计数增长；RP 降速；ECN 标记下降"),
        ("如果无效", "抓包/查 ACL/查交换机 CNP 计数器；确认没有中间设备重标记"),
    ]
    TABLE(doc, ["项", "内容"], rows, [3.0, 13.0], font_size=8.3, first_col_bold=True)

    H2(doc, "案例 5：总吞吐正常但公平性差", "10.5")
    rows = [
        ("观察", "总吞吐达标；Jain 低；部分流 P99 很差"),
        ("假设", "RTT 偏置、QP 数/优先级不均、PFC 反压不公平、交换机调度"),
        ("动作", "先按流/租户看分布；确认是否 RTT 差异；再评估 ECN/队列/调度"),
        ("验证", "Jain 改善；无流饿死；总吞吐不下降"),
        ("如果无效", "可能需要调度/整形/租户隔离，而不是继续调 DCQCN"),
    ]
    TABLE(doc, ["项", "内容"], rows, [3.0, 13.0], font_size=8.3, first_col_bold=True)

    H2(doc, "决策树", "10.6")
    CODE(doc, r"""
吞吐低?
├─ ECN 标记率高?
│  ├─ 是 → 队列高吗?
│  │     ├─ 是 → 拥塞严重：扩容/限流/查超配
│  │     └─ 否 → Kmin 太低/Pmax 太高：抬高 Kmin，必要时降 Pmax
│  └─ 否 → PFC pause 高?
│        ├─ 是 → PFC 在替 ECN：降低 Kmin，让 ECN 先动作
│        └─ 否 → 查 CNP/QoS/主机瓶颈
└─ 丢包?
   ├─ 是 → no-buffer-discard?
   │     ├─ 是 → headroom/cable/MTU/Xoff
   │     └─ 否 → 查链路/FEC/ACL
   └─ 否 → 查公平性/尾延迟/应用
""", size=8.2)

    H2(doc, "停止条件", "10.7")
    BULLETS(doc, [
        "lossless priority 出现持续 `no-buffer-discard`。",
        "PFC storm 计数增长，或交换机大面积 pause。",
        "CNP 计数异常下降，ECN 标记高但发送端不降速。",
        "业务 P99/吞吐超过 SLO，且 10 分钟内无法定位。",
        "重启后配置未恢复，且无法在窗口内修复。",
        "出现上述任一情况：停止调参，回滚，保留现场。",
    ])


def add_section_11(doc):
    H1(doc, "结论与检查清单", "11")
    P(doc, "拥塞调优的终点不是“参数看起来合理”，而是“在目标负载下可重复地满足 SLO，且没有引入新的风险”。")

    H2(doc, "五条结论", "11.1")
    NUMBERED(doc, [
        "**CX6 的 RoCEv2 主线是 DCQCN**；QCN 是 L2 遗留，ZTR-RTT/PCC 属于新平台/框架。[C1][C4][C7][C9]",
        "**ECN 是信号，CNP 是反馈，PFC 是兜底**；三者不能互相替代。[C3][C4][C11]",
        "**阈值要按时间理解**：Kmin/C 是队列时延；同样的字节数在不同速率下意义不同。",
        "**调优顺序**：正确性 → 缓冲 → ECN → CNP → DCQCN 恢复 → PFC 微调；一次只改一个变量。",
        "**验收要看分布和护栏**：P99、Jain、丢包、PFC storm，而不是平均吞吐。",
    ])

    H2(doc, "上线前检查清单", "11.2")
    checklist = [
        "已确认 CX6 端口是 Ethernet，RoCEv2 GID 正常。",
        "已记录固件/驱动/交换机 NOS/PSID/MTU/线缆长度。",
        "RoCE、CNP、普通流量使用不同 priority/TC；CNP strict。",
        "ECN 在 RoCE TC 启用；Kmin/Kmax/Pmax 单位与平台一致。",
        "PFC 只开 RoCE priority；两端一致；watchdog 已评估。",
        "headroom/cable length/Xoff/Xon 按最坏情况核对。",
        "DCQCN 参数使用默认基线；没有同时改多个参数。",
        "已做单流、4:1/8:1 incast、混合流量、故障/重启测试。",
        "已记录 P99、Jain、ECN 标记、PFC pause、no-buffer-discard。",
        "已准备回滚；没有用 mlxconfig reset 作为回滚手段。",
    ]
    for item in checklist:
        p = doc.add_paragraph()
        set_paragraph_format(p, before=1, after=2, line=1.25, left_indent=0.55)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        r = p.add_run("☐  ")
        set_run_font(r, BODY_FONT, 10.5, False, BLUE)
        add_rich_text(p, item, 10.0, NAVY)

    H2(doc, "下一步", "11.3")
    BULLETS(doc, [
        "用附录 B 的模板做一次完整基线，不要直接上生产调参。",
        "把第 8 章的实验矩阵落成可重复的脚本/流水线。",
        "把第 10 章的决策树接入值班手册。",
        "固件/驱动/交换机升级后，重跑基线和关键场景。",
    ])


def add_appendix_a(doc):
    H1(doc, "附录 A  公式与量纲", "A")
    P(doc, "以下公式用于量级判断和实验设计；标注“工程公式”的条目不是 NVIDIA 官方算法定义。")

    H2(doc, "A.1 常用公式", "A.1")
    CODE(doc, r"""
BDP（字节）        = C_Bps * RTT_s
队列时延（秒）      = q_bytes / C_Bps
Kmin 对应时延（秒） = Kmin_bytes / C_Bps
ECN 标记概率        = 0, q<Kmin
                     Pmax*(q-Kmin)/(Kmax-Kmin), Kmin<=q<Kmax
                     Pmax, q>=Kmax
Jain 公平性         = (sum(x_i))^2 / (n * sum(x_i^2))
DCQCN α 更新        = (1-g)*α + g  (有 CNP)
                     (1-g)*α      (无 CNP)
DCQCN 降速          = R*(1-α/2^RPG_GD)
DCQCN 恢复          = (R+Target)/2; Target += AI/HAI
""", size=8.0)

    H2(doc, "A.2 单位", "A.2")
    rows = [
        ("C_Bps", "字节/秒", "100GbE=12.5e9；200GbE=25e9"),
        ("RTT_s", "秒", "路径往返时间"),
        ("q_bytes", "字节", "Egress 队列占用"),
        ("Kmin/Kmax", "字节（或平台单位）", "Onyx absolute 用 KB；Cumulus 用 bytes"),
        ("α", "0–1", "DCQCN 拥塞评分"),
        ("RPG_GD", "指数", "默认 11；不是线性百分比"),
        ("Jain", "0–1", "1=完全公平"),
    ]
    TABLE(doc, ["符号", "单位", "说明"], rows, [2.6, 5.0, 8.4], font_size=8.3, first_col_bold=True)

    H2(doc, "A.3 量级示例", "A.3")
    P(doc, "100GbE、RTT=100 µs：BDP=12.5e9×100e-6=1.25 MB；150 KB 的 Kmin 约对应 12 µs 队列时延。200GbE 同样 150 KB 只对应约 6 µs，所以不能直接复用 100G 的阈值。")
    CALLOUT(doc, "公式的边界", "BDP/队列时延公式忽略内部 speedup、cell 开销、调度和 PFC 交互；用于判断量级，不用于精确配置。精确值以交换机厂商计算和实测为准。", "warn")


def add_appendix_b(doc):
    H1(doc, "附录 B  实验记录模板", "B")
    P(doc, "把每次调优当成一次实验记录；没有记录的“经验”，无法复用，也无法回滚。")

    H2(doc, "B.1 单次实验记录", "B.1")
    rows = [
        ("日期/人员", "谁在什么时间做的"),
        ("环境", "网卡/PSID/固件/驱动/交换机 NOS/端口速率/MTU/线缆"),
        ("配置", "priority、PFC、ECN TC、CNP DSCP/PCP、buffer、DCQCN 参数"),
        ("负载", "单流/incast 比例/流数/消息大小/持续时间"),
        ("变量", "只写本次改变的一个因素；原值/新值"),
        ("指标", "吞吐、P50/P99/P999、Jain、ECN 标记、PFC pause、丢包、重传"),
        ("结果", "与基线比较；是否达到效应量；护栏是否恶化"),
        ("结论", "保留/观察/回滚；下一步"),
    ]
    TABLE(doc, ["字段", "内容"], rows, [3.4, 12.6], font_size=8.3, first_col_bold=True)

    H2(doc, "B.2 测试矩阵", "B.2")
    rows = [
        ("单流", "1→1", "基线/链路", "3 次"),
        ("Incast 4:1", "4→1", "ECN/CNP/PFC", "3 次"),
        ("Incast 8:1", "8→1", "压力/公平", "3 次"),
        ("Permutation", "全互通", "公平/整体", "3 次"),
        ("Burst", "短时线速", "headroom", "3 次"),
        ("Mixed", "RoCE+TCP", "隔离", "3 次"),
        ("Reboot", "重启", "持久化", "1 次"),
    ]
    TABLE(doc, ["场景", "负载", "目的", "重复"], rows, [3.2, 4.2, 5.6, 3.0], font_size=8.3, first_col_bold=True)

    H2(doc, "B.3 验收记录", "B.3")
    BULLETS(doc, [
        "记录“通过/有条件通过/不通过”，不要只写“正常”。",
        "记录未达标项、影响范围、责任人和截止时间。",
        "记录回滚是否演练、耗时多少。",
        "把最终配置写入版本库/配置管理，不留在个人 shell 历史里。",
    ])


def add_appendix_c(doc):
    H1(doc, "附录 C  证据摘录", "C")
    P(doc, "下表摘录支撑关键结论的原文或原始论文表述，并给出中文释义。完整上下文见附录 D。")
    rows = [
        ("DCQCN 三角色", "\"The CP is typically the network switch... The NP is the receiving NIC... The RP is the sending NIC.\"", "CP=交换机，NP=接收网卡，RP=发送网卡。", "[C4]"),
        ("ECN 标记曲线", "\"As the queue size grows past Kmin, the probability of marking increases linearly until it reaches Pmax at Kmax.\"", "队列超过 Kmin 后标记概率线性上升，到 Kmax 达到 Pmax。", "[C4]"),
        ("CNP 生成", "\"When an ECN-marked packet arrives at the NP, a CNP is generated and sent back to the sender.\"", "NP 看到 CE 后生成 CNP 回送源端。", "[C4]"),
        ("α/降速/恢复", "\"alpha\" / \"rate is cut\" / \"Fast Recovery / Additive Increase / Hyper Increase\"", "官方文档分别描述 α 更新、乘法降速和三段恢复；这是关键词摘录，不是连续引文。", "[C4]"),
        ("CNP 默认 QoS", "\"CNP_DSCP ... 48\" / \"CNP_802P_PRIO ... 6\"", "CNP 默认 DSCP 48、PCP 6。", "[C5]"),
        ("ZTR-RTT 信号", "\"RTT\" + \"CNP\" / \"Additive Increase\" + \"Multiplicative Decrease\"", "ZTR-RTT 用 RTT+CNP，AIMD；关键词摘录。", "[C7]"),
        ("ZTR-RTT 平台", "\"within the DOCA Programmable Congestion Control (PCC) framework\"", "ZTR-RTT 属于新平台/框架，不应假定 CX6 支持。", "[C7]"),
        ("PCC 平台", "\"DOCA PCC ... BlueField-3 platform hardware acceleration\"", "PCC 依赖 BlueField-3/DPA。", "[C9]"),
        ("Onyx ECN 单位", "\"minimum-absolute ... (in KBs)\"", "Onyx absolute ECN 阈值单位是 KB。", "[C13]"),
        ("Cumulus ECN 默认", "\"minimum buffer threshold of 150000 bytes ... maximum 1500000 bytes\"", "Cumulus 默认 150 KB/1.5 MB（字节表示）。", "[C14]"),
        ("PFC watchdog", "\"PFC watchdog detects abnormal back pressure ... disables PFC temporarily\"", "watchdog 检测异常反压并临时止损。", "[C14]"),
        ("ECN 协议语义", "\"ECN uses two bits in the IP header\" / \"ECT/CE\"", "ECN 用 IP 头 2 bit，ECT/CE。", "[C19]"),
        ("DCQCN 论文", "\"Congestion Control for Large-Scale RDMA Deployments\"", "DCQCN 的算法来源与设计背景。", "[C18]"),
    ]
    TABLE(doc, ["结论", "原文/论文表述（节选）", "中文释义", "出处"], rows, [3.2, 8.4, 3.0, 1.4], font_size=7.6, first_col_bold=True)


def add_appendix_d(doc):
    H1(doc, "附录 D  来源与链接", "D")
    P(doc, "以下 21 份来源是本文的事实/算法依据，访问日期 2026-10-25。NVIDIA 文档可能随版本更新；链接失效时用标题检索。")

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, BORDER, 4)
    widths = [1.1, 7.0, 5.4, 2.5]
    for i, w in enumerate(widths):
        table.columns[i].width = Cm(w)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    headers = ["ID", "文档/论文（版本）", "用途", "链接"]
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_width(cell, widths[i])
        set_cell_background(cell, NAVY)
        set_cell_margins(cell, 0.07, 0.07, 0.10, 0.10)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        set_paragraph_format(p, before=0, after=0, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(h)
        set_run_font(r, HEAD_FONT, 8.6, True, WHITE)

    for idx, s in enumerate(TOPIC_SOURCES):
        tr = table.add_row()
        set_row_cant_split(tr)
        vals = [s["id"], s["title"] + "\n" + s["version"], s["use"], "打开"]
        for i, val in enumerate(vals):
            cell = tr.cells[i]
            set_cell_width(cell, widths[i])
            set_cell_margins(cell, 0.05, 0.05, 0.09, 0.09)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if idx % 2 == 1:
                set_cell_background(cell, "F7F9FC")
            p = cell.paragraphs[0]
            set_paragraph_format(p, before=0, after=0, line=1.15,
                                 align=WD_ALIGN_PARAGRAPH.CENTER if i in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT)
            if i == 3:
                add_hyperlink(p, s["url"], "官方页面", color=BLUE, size=7.8)
            else:
                for j, part in enumerate(str(val).split("\n")):
                    if j:
                        p.add_run().add_break()
                    r = p.add_run(part)
                    set_run_font(r, BODY_FONT, 7.8, i == 0, NAVY)

    H2(doc, "修订记录")
    rows = [
        ("v1.0", TOPIC_DATE, "初版：算法地图、ECN/DCQCN 深度讲解、参数敏感性、实验方法、分场景配方、案例与决策树。"),
    ]
    TABLE(doc, ["版本", "日期", "说明"], rows, [2.0, 3.0, 11.0], font_size=8.4, first_col_bold=True)

    H2(doc, "免责声明")
    P(doc, "本文基于公开 NVIDIA 文档、RFC 与 ACM 论文整理，未在读者的具体硬件/固件/业务上做实测；调优建议不构成性能承诺。生产变更前必须在同等环境验证，并准备回滚方案。")


def main():
    doc = Document()
    setup_styles(doc)
    add_cover(doc)
    add_front_matter(doc)
    add_body_start(doc)
    add_section_1(doc)
    add_section_2(doc)
    add_section_3(doc)
    add_section_4(doc)
    add_section_5(doc)
    add_section_6(doc)
    add_section_7(doc)
    add_section_8(doc)
    add_section_9(doc)
    add_section_10(doc)
    add_section_11(doc)
    add_appendix_a(doc)
    add_appendix_b(doc)
    add_appendix_c(doc)
    add_appendix_d(doc)

    doc.core_properties.title = TOPIC_TITLE
    doc.core_properties.subject = TOPIC_SUBTITLE
    doc.core_properties.author = "网络技术整理"
    doc.core_properties.keywords = "DCQCN, ECN, PFC, ZTR-RTT, PCC, RoCEv2, congestion control tuning"
    doc.core_properties.comments = "Companion to the ConnectX-6 RoCEv2 CC/PFC guide v1.1."
    doc.save(TOPIC_DOCX)
    print("saved", TOPIC_DOCX)


if __name__ == "__main__":
    main()
