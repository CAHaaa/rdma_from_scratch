# -*- coding: utf-8 -*-
"""Cover, front matter, and chapters 1-5 of the CC/tuning companion."""
import os
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docx_helpers import (
    NAVY, BLUE, TEAL, ORANGE, RED, GRAY, WHITE,
    LIGHT_BLUE, LIGHT_GRAY, LIGHT_GREEN, LIGHT_ORANGE, LIGHT_RED,
    BODY_FONT, HEAD_FONT,
    P, BULLETS, NUMBERED, H1, H2, H3, CAPTION, TABLE, CODE, CALLOUT, FIGURE,
    set_run_font, set_paragraph_format, add_rich_text,
    set_cell_background, set_cell_margins, set_cell_width,
)
from common_topic import TOPIC_TITLE, TOPIC_SUBTITLE, TOPIC_VERSION, TOPIC_DATE


def add_cover(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.1)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)
    sec.different_first_page_header_footer = True

    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell = t.cell(0, 0)
    set_cell_width(cell, 16.2)
    set_cell_background(cell, NAVY)
    set_cell_margins(cell, 0.28, 0.28, 0.30, 0.30)
    p = cell.paragraphs[0]
    set_paragraph_format(p, before=0, after=0, line=1.0)
    r = p.add_run("NVIDIA NETWORKING   ·   DEEP DIVE")
    set_run_font(r, HEAD_FONT, 9.5, True, "D9E7F5")

    p = doc.add_paragraph()
    set_paragraph_format(p, before=20, after=4, line=1.12)
    r = p.add_run("ConnectX-6 RoCEv2")
    set_run_font(r, HEAD_FONT, 24, True, NAVY)

    p = doc.add_paragraph()
    set_paragraph_format(p, before=0, after=4, line=1.15)
    r = p.add_run("拥塞算法与调优专题")
    set_run_font(r, HEAD_FONT, 24, True, NAVY)

    p = doc.add_paragraph()
    set_paragraph_format(p, before=0, after=12, line=1.25)
    r = p.add_run(TOPIC_SUBTITLE)
    set_run_font(r, BODY_FONT, 11.5, False, GRAY)

    p = doc.add_paragraph()
    set_paragraph_format(p, before=4, after=14, line=1.0)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), TEAL)
    pBdr.append(bottom)
    pPr.append(pBdr)

    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    c = t.cell(0, 0)
    set_cell_width(c, 15.8)
    set_cell_background(c, LIGHT_BLUE)
    set_cell_margins(c, 0.18, 0.18, 0.22, 0.22)
    p = c.paragraphs[0]
    set_paragraph_format(p, before=0, after=0, line=1.2)
    r = p.add_run("算法原理  ·  参数敏感性  ·  实验设计  ·  分场景调优  ·  21 份依据")
    set_run_font(r, BODY_FONT, 9.4, True, BLUE)

    p = doc.add_paragraph()
    set_paragraph_format(p, before=20, after=6, line=1.2)
    r = p.add_run("与主指南的关系")
    set_run_font(r, HEAD_FONT, 10.0, True, TEAL)
    p = doc.add_paragraph()
    set_paragraph_format(p, before=0, after=0, line=1.45)
    r = p.add_run("配套《ConnectX-6 网卡 RoCEv2 拥塞控制与 PFC 流控实施指南 v1.1》。\n主指南回答“怎么配”，本专题回答“为什么这样调、如何证明调对了”。")
    set_run_font(r, BODY_FONT, 9.8, False, NAVY)

    p = doc.add_paragraph()
    set_paragraph_format(p, before=14, after=6, line=1.2)
    r = p.add_run("适用对象")
    set_run_font(r, HEAD_FONT, 10.0, True, TEAL)
    p = doc.add_paragraph()
    set_paragraph_format(p, before=0, after=0, line=1.45)
    r = p.add_run("网络架构师、RDMA 性能工程师、AI/HPC 集群调优与测试人员")
    set_run_font(r, BODY_FONT, 9.8, False, NAVY)

    doc.add_paragraph().paragraph_format.space_after = Pt(24)
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    vals = [("版本", TOPIC_VERSION), ("日期", TOPIC_DATE), ("文档状态", "技术专题稿 · 与主指南配套")]
    for i, (k, v) in enumerate(vals):
        c = t.cell(0, i)
        set_cell_width(c, 5.4)
        set_cell_margins(c, 0.10, 0.10, 0.12, 0.12)
        p = c.paragraphs[0]
        set_paragraph_format(p, before=0, after=0, line=1.2, align=WD_ALIGN_PARAGRAPH.LEFT)
        r = p.add_run(k + "\n")
        set_run_font(r, BODY_FONT, 8.2, False, GRAY)
        r = p.add_run(v)
        set_run_font(r, BODY_FONT, 9.0, True, NAVY)

    p = doc.add_paragraph()
    set_paragraph_format(p, before=30, after=0, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("本文件为公开资料的工程整理；算法结论以 NVIDIA 文档与原始论文为准，调优建议需现场验证。")
    set_run_font(r, BODY_FONT, 8.0, False, GRAY)


def add_front_matter(doc):
    doc.add_page_break()
    H2(doc, "文档信息")
    rows = [
        ("文档名称", TOPIC_TITLE),
        ("副标题", TOPIC_SUBTITLE),
        ("版本 / 日期", f"{TOPIC_VERSION} / {TOPIC_DATE}"),
        ("证据基线", "21 份来源：NVIDIA 官方文档、DOCA/MFT 文档、RFC、ACM 论文"),
        ("适用平台", "ConnectX-6 EN/VPI 的 Ethernet 模式；RoCEv2；Linux"),
        ("配套文档", "《ConnectX-6 网卡 RoCEv2 拥塞控制与 PFC 流控实施指南 v1.1》"),
    ]
    TABLE(doc, ["项目", "内容"], rows, [3.2, 13.0], font_size=8.7, first_col_bold=True)

    H2(doc, "与主指南的分工")
    rows = [
        ("主指南", "机制定位、端到端设计、命令、验收、排障、回滚", "回答“怎么配、怎么验”"),
        ("本专题", "算法原理、参数敏感性、实验设计、分场景调优配方", "回答“为什么这样调、如何调得更稳”"),
        ("共享结论", "PFC 是逐跳流控；DCQCN 是端到端 CC；CX6 主线是 ECN+CNP+DCQCN", "两份文档保持一致"),
    ]
    TABLE(doc, ["文档", "重点", "回答的问题"], rows, [3.0, 9.2, 4.0], font_size=8.5, first_col_bold=True)

    H2(doc, "阅读说明与证据等级")
    P(doc, "正文用 [C1]–[C21] 标注来源。NVIDIA 文档用于能力、默认值、命令和平台边界；RFC 用于协议语义；ACM 论文用于算法来源与设计背景。工程建议和工程公式会单独标注，不与官方事实混淆。")
    rows = [
        ("官方事实", "NVIDIA 产品文档中的能力、默认值、命令、限制", "以 [C#] 标注"),
        ("标准/论文", "RFC 的协议语义、ACM 论文的算法设计", "以 [C#] 标注"),
        ("工程建议", "调优顺序、实验方法、经验阈值", "文中写“建议/通常/起始假设”"),
        ("工程公式", "BDP、队列时延、headroom、Jain 指数等", "文中标注“工程公式”"),
    ]
    TABLE(doc, ["等级", "含义", "标记"], rows, [3.0, 9.8, 3.4], font_size=8.5, first_col_bold=True)

    H2(doc, "目录")
    toc = [
        ("1", "专题导读：从“能跑”到“跑稳”", "7", "参数敏感性与交互"),
        ("2", "拥塞控制基础：指标、时间尺度与闭环", "8", "调优方法论：基线、实验与统计"),
        ("3", "算法地图：DCQCN、QCN、ZTR-RTT、PCC", "9", "分场景调优配方"),
        ("4", "ECN / RED / WRED：交换机侧信号", "10", "调优案例与决策树"),
        ("5", "DCQCN 深度讲解", "11", "结论与检查清单"),
        ("6", "其他算法：QCN、ZTR-RTT、PCC", "A", "附录 A  公式与量纲"),
        ("B", "附录 B  实验记录模板", "C", "附录 C  证据摘录"),
        ("D", "附录 D  来源与链接", "", ""),
    ]
    TABLE(doc, ["章", "内容", "章", "内容"], toc, [1.2, 6.8, 1.2, 6.8], font_size=8.2, first_col_bold=True, align_center=False)


def add_section_1(doc):
    h = H1(doc, "专题导读：从“能跑”到“跑稳”", "1")
    h.paragraph_format.page_break_before = False
    P(doc, "“能跑”只需要链路通、RoCE 能建链；“跑稳”要求在 incast、突发、故障、混合流量下，吞吐、时延、公平性和丢包都可控。两者之间差的不是某一个参数，而是一套可复现的调优方法。")

    H2(doc, "为什么默认值不够", "1.1")
    BULLETS(doc, [
        "NVIDIA 默认值是通用起点，不是为你的 incast 比例、MTU、RTT、业务突发模式选的。[C4][C5]",
        "同一个 ECN 阈值，在 100G/100m 和 200G/2km 上对应的队列时延完全不同；只抄数字没有意义。",
        "PFC 和 ECN 作用在不同缓冲：PFC 看入口 PG，ECN 看出口队列；不能用一个阈值去“对齐”另一个。",
        "业务目标不同：AI 训练怕长尾和同步突发，存储怕丢包和重传，多租户怕相互影响；调优方向不可能一样。",
        "默认值 + 错误优先级（比如 CNP 没走 strict）时，调任何 CC 参数都可能无效。",
    ])

    H2(doc, "本专题要建立的三个能力", "1.2")
    rows = [
        ("看懂算法", "知道 DCQCN 为什么用 ECN/CNP、α 和 RPG_GD 在做什么、ZTR-RTT/PCC 为什么不是 CX6 的默认路径", "第 2–6 章"),
        ("看懂参数", "知道每个参数调大/调小会改变什么、和谁耦合、先调谁", "第 7 章"),
        ("看懂实验", "能用单流、incast、突发、混合流量证明“调对了”，而不是只看峰值吞吐", "第 8–10 章"),
    ]
    TABLE(doc, ["能力", "目标", "章节"], rows, [3.2, 9.8, 3.0], font_size=8.5, first_col_bold=True)

    H2(doc, "三条调优原则", "1.3")
    NUMBERED(doc, [
        "**先正确，再优化**：优先级、CNP QoS、PFC 域、buffer/headroom 不正确时，不要调 DCQCN 参数。",
        "**一次只改一个变量**：同时改 Kmin、RPG_GD、AI/HAI，出了问题无法归因。",
        "**用分布，不用平均值**：看 P50/P99/P999、ECN 标记率、PFC pause、Jain 公平性，不只看平均吞吐。",
    ])
    CALLOUT(doc, "调优不是“找一组神参数”", "DCQCN 的稳定性取决于交换机标记、CNP 回送、端点速率控制和业务负载的共同作用。把某组参数从别人的集群抄过来，往往比默认值更差。", "warn")

    H2(doc, "一页结论", "1.4")
    rows = [
        ("算法", "CX6 的 RoCEv2 主线是 DCQCN；QCN 是 L2 遗留；ZTR-RTT/PCC 属于更新平台/框架。", "[C1][C4][C7][C9]"),
        ("信号", "ECN 是早期信号，CNP 是反馈，PFC 是逐跳兜底；三者不能互相替代。", "[C3][C4][C11]"),
        ("调优顺序", "正确性 → 缓冲/headroom → ECN 阈值 → CNP QoS → DCQCN 恢复参数 → PFC 微调。", "工程建议"),
        ("验证", "单流、4:1/8:1 incast、突发、混合流量、重启持久化，至少五类。", "工程建议"),
        ("安全", "每次变更都要有基线、回滚、停止条件；不要在不知道原值的情况下改参数。", "工程建议"),
    ]
    TABLE(doc, ["主题", "结论", "依据"], rows, [3.0, 10.6, 2.4], font_size=8.4, first_col_bold=True)


def add_section_2(doc):
    H1(doc, "拥塞控制基础：指标、时间尺度与闭环", "2")
    P(doc, "拥塞控制不是“把队列压到 0”。它是在吞吐、时延、公平性、丢包之间做权衡。理解这个权衡，才能判断一个参数该不该动。")

    H2(doc, "拥塞从哪里来", "2.1")
    rows = [
        ("持续超配", "下行带宽小于上行汇聚带宽", "队列持续增长，ECN/PFC 长期生效", "扩容/限速/流量工程"),
        ("Incast", "多个发送端同时打一个接收端/一个出口", "微秒级队列尖峰，平均利用率可能不高", "ECN 早期标记、PFC 兜底、应用限流"),
        ("Microburst", "短时间线速突发，平均速率不高", "PFC 触发但 ECN 不一定标记", "headroom、入口缓冲、 pacing"),
        ("速度不匹配", "100G 进 25G 出、或交换机内部/外部速率差", "出口队列堆积", "QoS/整形/端口规划"),
        ("慢接收端", "应用不 drain、QP 饥饿、CPU/GPU 瓶颈", "PFC 反压扩散，甚至风暴", "修应用/接收端，不是调 CC"),
        ("多租户/混合流量", "TCP 与 RoCE 共享 buffer/priority", "互相挤占，PFC 误伤", "优先级隔离、独立 lossless pool"),
    ]
    TABLE(doc, ["来源", "典型场景", "表现", "优先动作"], rows, [3.0, 5.4, 4.8, 2.8], font_size=8.2, first_col_bold=True)

    H2(doc, "控制回路：谁在控制谁", "2.2")
    P(doc, "把拥塞控制看成闭环：交换机队列是被控对象，ECN 是传感器，CNP 是反馈通道，发送端限速器是执行器。调参本质上是在改控制器增益和反馈延迟。")
    rows = [
        ("被控对象", "交换机 Egress 队列 / 入口 PG", "队列长度、占用率、丢弃", "buffer、阈值、调度"),
        ("传感器", "ECN CE 标记", "标记概率、标记率", "Kmin/Kmax/Pmax"),
        ("反馈", "CNP", "CNP 速率、时延、丢失", "CNP DSCP/PCP、strict、最小间隔"),
        ("执行器", "发送端 RP 限速器", "CurrentRate/TargetRate", "α、RPG_GD、AI/HAI、定时器"),
        ("兜底", "PFC", "Pause 帧、pause duration", "Xoff/Xon、headroom、watchdog"),
    ]
    TABLE(doc, ["角色", "位置", "观测量", "可调项"], rows, [3.0, 5.0, 4.6, 3.4], font_size=8.4, first_col_bold=True)

    H2(doc, "时间尺度", "2.3")
    rows = [
        ("PFC 反应", "微秒级", "入口缓冲到 Xoff，发 Pause；链路两端 MAC/交换芯片处理", "PFC 是最后一道防线"),
        ("ECN 标记", "微秒–毫秒", "Egress 队列超过 Kmin，按概率标记", "标记越早，越早给端点信号"),
        ("CNP 反馈", "RTT 量级", "接收端生成 CNP，回到发送端", "CNP 延迟/丢失会直接拖慢反应"),
        ("α 更新", "DCE_TCP_RTT（默认 1 µs）", "有/无 CNP 的 EWMA", "默认极快，不要按“RTT”字面理解"),
        ("降速窗口", "RATE_REDUCE_MONITOR_PERIOD（默认 4 µs）", "观察 CNP 后降速", "窗口越短越敏感"),
        ("恢复", "RPG_TIME_RESET 300 µs / BYTE_RESET ~2 MB", "AI/HAI 逐步升速", "恢复过快会再次拥塞"),
    ]
    TABLE(doc, ["阶段", "时间尺度", "发生什么", "调优含义"], rows, [3.4, 4.4, 5.2, 3.0], font_size=8.1, first_col_bold=True)
    CALLOUT(doc, "DCE_TCP_RTT 不是实测 RTT", "DOCA 文档给的单位是 µs、默认 1；它是 α 更新周期，不是路径 RTT 测量。把它当成“RTT”会误判调优方向。[C5]", "warn")

    H2(doc, "指标：不要只看平均吞吐", "2.4")
    rows = [
        ("吞吐/Goodput", "完成的有用数据速率", "平均 + 分位", "平均值掩盖长尾"),
        ("时延", "P50/P99/P999、队列时延", "分位 + 时间序列", "P99 才能看到 incast/排队"),
        ("丢包", "`no-buffer-discard`、重传、out-of-sequence", "计数 + 速率", "lossless 下应为 0 或接近 0"),
        ("ECN", "标记率、标记包数", "按 TC/端口", "标记率高不一定是坏事，要看队列"),
        ("PFC", "pause 包数/时长、storm 计数", "按 priority", "高 pause + 低 ECN = PFC 在替 ECN 干活"),
        ("公平性", "Jain 指数、各流吞吐分布", "多流/多租户", "平均值相同也可能极不公平"),
        ("收敛/稳定", "恢复时间、振荡幅度", "阶跃/incast 后", "只看稳态会漏掉振荡"),
    ]
    TABLE(doc, ["指标", "定义/来源", "维度", "注意"], rows, [3.2, 6.0, 3.4, 3.2], font_size=8.2, first_col_bold=True)
    P(doc, "Jain 公平性指数（工程公式）：`J = (Σx_i)^2 / (n · Σx_i^2)`。J=1 表示完全公平；J 越低越不公平。它不替代业务 SLO，但能揭示“总吞吐正常、个别流饿死”的问题。")

    H2(doc, "为什么 RoCE 用 ECN 而不是丢包", "2.5")
    BULLETS(doc, [
        "RoCE 的丢包恢复依赖重传/乱序处理，代价比 TCP 更高；lossless 设计希望“提前降速”而不是“丢了再补”。[C4]",
        "ECN 在 IP 头标记 CE，不丢包；接收端网卡生成 CNP，发送端按 QP 降速，形成闭环。[C4][C19]",
        "PFC 是逐跳暂停，不告诉发送端“降多少”；它只能兜底，不能替代 CC。[C3][C11]",
        "因此 RoCEv2 的首选信号是 ECN，PFC 是最后一道防线；两者职责不同。",
    ])


def add_section_3(doc):
    H1(doc, "算法地图：DCQCN、QCN、ZTR-RTT、PCC", "3")
    P(doc, "先把“算法”与“机制”分开：PFC 是流控，QCN/DCQCN/ZTR-RTT 是拥塞控制，PCC 是可编程框架。把它们混在一起，就会得出“开 PFC 就能替代 CC”这类错误结论。")

    H2(doc, "一张图", "3.1")
    FIGURE(doc, os.path.join("assets", "fig_cc_map.png"), 16.0)
    CAPTION(doc, "图 1  RoCE 拥塞控制算法地图（按 [C1][C4][C7][C9][C11] 整理，概念图）")

    H2(doc, "算法家族", "3.2")
    rows = [
        ("丢包型", "TCP Reno 等", "以丢包作为拥塞信号，AIMD/慢启动", "通用互联网/TCP；RoCE 不希望丢包，恢复代价高", "[C21]"),
        ("ECN 型", "DCTCP、DCQCN", "ECN 标记 + 端点降速；DCQCN 用 CNP 反馈", "RoCEv2/数据中心；CX6 主线，需交换机支持 ECN", "[C4][C5]"),
        ("延迟型", "TIMELY、ZTR-RTT", "RTT/排队时延作为信号；AIMD/窗口", "数据中心/新平台；ZTR 在新平台，TIMELY 为研究方案", "[C7][C20]"),
        ("L2 CNM", "QCN（802.1Qau）", "交换机生成 CNM，端点降速", "Legacy L2/RoCEv1；L2、不可路由", "[C1]"),
        ("信用", "IB Credit", "链路信用，逐跳保证不溢出", "InfiniBand；不是 Ethernet PFC/ECN", "[C1]"),
        ("可编程", "DOCA PCC", "用户实现 CC 算法，DPA 执行", "BlueField-3/新平台；开发门槛高，非 CX6 即用", "[C9]"),
    ]
    TABLE(doc, ["家族", "代表", "核心信号/动作", "适用与边界", "依据"], rows,
          [2.2, 3.2, 5.2, 4.0, 1.4], font_size=7.8, first_col_bold=True)

    H2(doc, "CX6 上应该用哪一个", "3.3")
    rows = [
        ("DCQCN", "CX6 的 RoCEv2 主线；ECN + CNP + 速率控制", "默认选它；先保证 ECN/CNP QoS 正确", "[C1][C4][C5]"),
        ("QCN", "CX6 手册列有 802.1Qau 能力；L2 遗留", "只有既有 L2/QCN 环境才考虑；不要与 DCQCN 混用", "[C1]"),
        ("ZTR-RTT", "DOCA 3.5 的新一代原生算法；RTT+CNP", "不要假设 CX6 支持；按型号/固件文档确认", "[C6][C7]"),
        ("PCC", "可编程 CC；BlueField-3 + DPA", "CX6 不按 PCC 方式配置；需要新平台/开发投入", "[C9]"),
        ("TIMELY", "RTT 型研究方案", "用于理解延迟型 CC，不作为 CX6 配置项", "[C20]"),
    ]
    TABLE(doc, ["算法", "定位", "在 CX6 上的用法", "依据"], rows, [2.8, 7.0, 4.2, 2.0], font_size=8.0, first_col_bold=True)

    H2(doc, "选择算法时的四个问题", "3.4")
    NUMBERED(doc, [
        "业务能否容忍丢包/重传？不能 → lossless + ECN/PFC；能 → lossy + ECN/重传。",
        "路径是否可路由/跨 L3？是 → ECN/DCQCN 或 ZTR-RTT；否 → 可考虑 L2/QCN，但优先 DCQCN。",
        "交换机是否支持 ECN/CNP QoS？否 → 先补网络能力，或评估 ZTR-RTT 平台；不要把 DCQCN 硬开在无标记路径上。",
        "是否愿意/能够开发自定义算法？是 → PCC（新平台）；否 → 用原生 DCQCN/ZTR-RTT。",
    ])
    CALLOUT(doc, "不要把“算法”当“配置开关”", "DCQCN 不是网卡上的一个 enable 位；它依赖交换机 ECN 标记、CNP 的 QoS、端点参数和缓冲设计。只开 `ROCE_CC_PRIO_MASK` 而不配 ECN/CNP，等于闭环断了两段。", "risk")


def add_section_4(doc):
    H1(doc, "ECN / RED / WRED：交换机侧信号", "4")
    P(doc, "ECN 是 DCQCN 的“传感器”。交换机只负责在队列增长时把 ECT 包标记为 CE；它不直接降速，降速由端点完成。把 ECN 配成“要么不标记、要么全标记”，闭环都会失效。")

    H2(doc, "ECN 的协议语义", "4.1")
    P(doc, "ECN 用 IP 头 2 bit 表示：00 Not-ECT、01/10 ECT、11 CE。支持 ECN 的交换机在队列超过阈值时把 ECT 改成 CE；接收端看到 CE 后生成 CNP。它要求路径上所有节点支持 ECN，否则标记无法端到端传递。[C2][C19]")
    rows = [
        ("00", "Not-ECT", "不支持 ECN；拥塞时通常丢弃"),
        ("01", "ECT(1)", "支持 ECN，可被标记为 CE"),
        ("10", "ECT(0)", "支持 ECN，可被标记为 CE"),
        ("11", "CE", "已遇到拥塞；接收端生成 CNP"),
    ]
    TABLE(doc, ["ECN 位", "名称", "含义"], rows, [2.4, 4.0, 9.8], font_size=8.6, first_col_bold=True)

    H2(doc, "标记曲线", "4.2")
    FIGURE(doc, os.path.join("assets", "fig_ecn_curve.png"), 16.0)
    CAPTION(doc, "图 2  ECN 标记概率与 Kmin/Kmax/Pmax 的关系（概念图，依据 [C4][C13]）")
    P(doc, "DCQCN 的拥塞点通常是交换机 Egress 队列：q < Kmin 不标记；Kmin ≤ q < Kmax 线性增加；q ≥ Kmax 达到 Pmax。Cumulus 的默认 RoCE profile 把 Pmax 设为 100%，即在 Kmax 时所有 ECN-capable 包都被标记。[C4][C14]")
    CODE(doc, r"""
q < Kmin:                 p = 0
Kmin <= q < Kmax:         p = Pmax * (q - Kmin) / (Kmax - Kmin)
q >= Kmax:                p = Pmax

工程换算：
队列时延 qdelay ≈ q / C
Kmin 对应的时延 ≈ Kmin / C
""", size=8.0)

    H2(doc, "RED、WRED 与 ECN 的区别", "4.3")
    rows = [
        ("RED", "随机早期检测", "队列超过阈值时按概率丢弃", "有损流量/TCP", "用丢包通知拥塞"),
        ("WRED", "加权 RED", "按颜色/优先级使用不同阈值", "DiffServ/多优先级", "更细的丢弃策略"),
        ("ECN", "显式拥塞通知", "把 ECT 标记为 CE，不丢弃", "RoCEv2/DCQCN", "用标记通知拥塞"),
    ]
    TABLE(doc, ["机制", "全称", "动作", "典型用途", "信号"], rows, [2.2, 3.2, 4.8, 3.4, 2.4], font_size=8.2, first_col_bold=True)
    P(doc, "对 lossless RoCE，优先 ECN；RED 可作为有损流量的补充，但不要让 RED 先于 ECN 丢掉 RoCE 包。Onyx 支持 `traffic-class <tc> congestion-control [red|ecn|both]`，阈值可用绝对 KB 或相对百分比；两者不要混着配。[C13]")

    H2(doc, "Onyx 与 Cumulus 的配置差异", "4.4")
    rows = [
        ("Onyx", "`traffic-class <tc> congestion-control ecn minimum-absolute <KB> maximum-absolute <KB>`", "绝对阈值单位 KB；也支持 relative 百分比", "[C13][C10]"),
        ("Cumulus", "`nv set qos congestion-control default-global traffic-class <tc> min-threshold <bytes> max-threshold <bytes>`", "默认 profile 以字节表示；RoCE profile 用 TC3", "[C14][C15]"),
        ("共同点", "Kmin/Kmax/Pmax 语义一致；都要按 TC、按端口核对", "不要把 Onyx 的 150 当字节，也不要把 Cumulus 的 150000 当 KB", "[C13][C14]"),
    ]
    TABLE(doc, ["平台", "配置入口", "单位/注意", "依据"], rows, [2.6, 8.2, 3.6, 1.6], font_size=8.0, first_col_bold=True)

    H2(doc, "从“队列时延”理解阈值", "4.5")
    P(doc, "把 Kmin 换算成时间，比记数字更有意义：`Kmin_time ≈ Kmin_bytes / C`。同样 150 KB，在 100G 上约 12 µs，在 200G 上约 6 µs；同样的阈值在不同速率下代表不同的排队时延。")
    rows = [
        ("100G", "12.5 GB/s", "1.25 MB", "150 KB ≈ 12 µs"),
        ("200G", "25 GB/s", "2.5 MB", "150 KB ≈ 6 µs"),
        ("400G", "50 GB/s", "5 MB", "150 KB ≈ 3 µs"),
    ]
    TABLE(doc, ["线速", "字节速率", "100 µs RTT 的 BDP", "150 KB 对应队列时延"], rows, [3.0, 4.2, 5.0, 3.8], font_size=8.4, first_col_bold=True)
    P(doc, "工程含义：不要问“Kmin 应该是多少字节”，先问“我希望 ECN 在多少 µs 队列时延开始标记”。目标时延由业务 SLO、RTT、交换机缓冲和 PFC 阈值共同决定。")

    H2(doc, "常见错误", "4.6")
    BULLETS(doc, [
        "把 Onyx 的 150 当字节，导致 ECN 几乎不标记；或把 Cumulus 的 150000 当 KB，导致过早全标记。",
        "只开 ECN、不开 CNP 的 strict QoS；标记了但反馈回不来。",
        "Kmin 设得比 PFC Xoff 还晚，PFC 先触发，DCQCN 看不到持续拥塞。",
        "Pmax 设得很低，队列已经很高却只标记少量包，端点降速不足。",
        "用 RED 替 ECN 丢 RoCE 包，再指望 DCQCN 恢复；方向反了。",
    ])


def add_section_5(doc):
    H1(doc, "DCQCN 深度讲解", "5")
    P(doc, "DCQCN 是 CX6 上 RoCEv2 的主线算法。它把 DCTCP 的 ECN 拥塞评分和 QCN 的速率控制引擎结合起来：用 ECN 感知拥塞，用 CNP 回送反馈，用 α 和乘法降速/加法恢复调节每个 QP 的速率。[C4][C18]")

    H2(doc, "三角色与闭环", "5.1")
    NUMBERED(doc, [
        "**CP（拥塞点）**：交换机 Egress 队列按 Kmin/Kmax/Pmax 标记 CE。[C4]",
        "**NP（通知点）**：接收端网卡看到 CE，生成 CNP，按 QP 回送；硬件完成，不经过 CPU。[C4]",
        "**RP（反应点）**：发送端网卡收到 CNP，更新 α，降速；没有 CNP 时逐步恢复。[C4]",
        "**闭环**：降速 → 队列下降 → 停止标记 → CNP 减少 → 升速；再次拥塞则重复。",
    ])
    P(doc, "CNP 是按流/QP 反馈的：一个 QP 拥塞，不会让同网卡所有 QP 一起降速。这是 DCQCN 相对全局 pause 的关键优势，也是它必须和 PFC 分开理解的原因。[C4]")

    H2(doc, "状态机", "5.2")
    FIGURE(doc, os.path.join("assets", "fig_dcqcn_state.png"), 16.0)
    CAPTION(doc, "图 3  DCQCN 的降速、Fast Recovery、AI、HAI 与 α 更新（依据 [C4][C5] 绘制）")
    P(doc, "收到 CNP 后先做乘法降速，再进入 Fast Recovery；连续无 CNP 的触发达到阈值后进入 AI，再达到双计数器阈值后进入 HAI。α 表示近期拥塞程度，随 CNP 出现而升高、随静默而衰减。[C4]")

    H2(doc, "公式", "5.3")
    CODE(doc, r"""
alpha 更新（每个 DCE_TCP_RTT）:
  有 CNP:   alpha <- (1 - g) * alpha + g
  无 CNP:   alpha <- (1 - g) * alpha
  g = DCE_TCP_G / 1024

降速（收到 CNP，按 RATE_REDUCE_MONITOR_PERIOD 观察）:
  CurrentRate <- CurrentRate * (1 - alpha / 2^RPG_GD)

恢复（无 CNP）:
  CurrentRate <- (CurrentRate + TargetRate) / 2
  TargetRate  <- TargetRate + AI    (Additive Increase)
  TargetRate  <- TargetRate + HAI   (Hyper Increase)
""", size=8.2)
    P(doc, "默认 `DCE_TCP_G=1019`、`RPG_GD=11`。按公式直算，α=1 时单次降速约 1/2048；所以 DCQCN 的默认行为是“小步快跑”，依赖持续的 CNP 和 4 µs 的监控窗口，而不是一次 CNP 砍半。把 `RPG_GD` 调小会让单次降速更狠，调大更温和。[C5]")

    H2(doc, "参数与行为", "5.4")
    rows = [
        ("ROCE_CC_PRIO_MASK", "Enable（逐优先级）", "按 priority 启用 DCQCN；255=全开", "默认优先 [C5][C17]"),
        ("DCE_TCP_G", "1019", "α 的 EWMA 权重", "默认优先 [C5]"),
        ("DCE_TCP_RTT", "1 µs", "α 更新周期", "默认优先 [C5]"),
        ("INITIAL_ALPHA_VALUE", "1023", "首个 CNP 的 α", "默认优先 [C5]"),
        ("RATE_TO_SET_ON_FIRST_CNP", "0", "首个 CNP 是否强制速率", "默认优先 [C5]"),
        ("RPG_GD", "11", "降速指数；越小越狠", "谨慎 [C5]"),
        ("RPG_MIN_DEC_FAC", "50%", "单次降速上限", "谨慎 [C5]"),
        ("RPG_MIN_RATE", "1 Mb/s", "QP 最低速率", "谨慎 [C5]"),
        ("RATE_REDUCE_MONITOR_PERIOD", "4 µs", "CNP 观察窗口", "谨慎 [C5]"),
        ("RPG_TIME_RESET", "300 µs", "升速定时器", "谨慎 [C5]"),
        ("RPG_BYTE_RESET", "32767×64 B≈2 MB", "升速字节计数器", "谨慎 [C5]"),
        ("RPG_THRESHOLD", "1", "AI/HAI 的连续触发次数", "谨慎 [C5]"),
        ("RPG_AI_RATE / HAI_RATE", "5 / 50 Mb/s", "AI/HAI 步长", "谨慎 [C5]"),
        ("CNP_DSCP / CNP_802P_PRIO", "48 / 6", "CNP 的 DSCP/PCP", "先保证 QoS [C5]"),
        ("MIN_TIME_BETWEEN_CNPS", "4 µs", "CNP 最小间隔", "谨慎 [C5]"),
    ]
    TABLE(doc, ["参数", "默认", "作用", "调优优先级"], rows, [5.2, 2.8, 5.6, 2.4], font_size=7.9, first_col_bold=True)

    H2(doc, "DCQCN 的公平性与收敛", "5.5")
    BULLETS(doc, [
        "**按 QP 反馈**：每个 QP 独立降速；天然支持多流，但也意味着 QP 数越多，控制面越复杂。",
        "**RTT 偏置**：长 RTT 流收到 CNP 更晚、恢复更慢，可能出现不公平；DCQCN 本身不保证 RTT 公平。",
        "**α 的记忆**：默认 g 很大，α 对单个 CNP 很敏感、对静默衰减也快；这使系统反应快，但也更容易受 CNP 到达模式影响。",
        "**收敛不是无条件**：交换机阈值、CNP 速率、QP 数、RTT、PFC 交互都会影响收敛；必须用 incast 实测，不能只看公式。",
        "**PFC 会掩盖拥塞**：如果 PFC 先触发，发送端可能收不到足够 CNP，速率不降，PFC 变成主要反压；这是“PFC 很忙、ECN 很闲”的根因之一。",
    ])

    H2(doc, "DCQCN vs DCTCP vs QCN", "5.6")
    rows = [
        ("DCTCP", "TCP/ECN", "ACK 携带 ECN-Echo", "软件 TCP", "数据中心 TCP"),
        ("DCQCN", "RoCEv2/ECN", "CNP 硬件回送", "网卡硬件", "CX6 RoCE 主线"),
        ("QCN", "L2/802.1Qau", "CNM", "网卡/交换机", "Legacy L2"),
    ]
    TABLE(doc, ["算法", "层/信号", "反馈", "执行", "定位"], rows, [2.4, 4.0, 4.4, 3.0, 2.2], font_size=8.3, first_col_bold=True)
    P(doc, "DCQCN 不是“更高级的 QCN”，而是为可路由 RoCEv2 重新设计的 ECN/CNP 闭环；QCN 的 L2 CNM 不能跨路由域，不能拿来替代 DCQCN。[C1][C4]")

    H2(doc, "DCQCN 做不到什么", "5.7")
    BULLETS(doc, [
        "不能创造带宽；持续超配/严重 incast 需要容量/流量工程。",
        "不能替代 PFC；严重瞬时溢出仍需入口缓冲/PFC 兜底。",
        "不能感知不标记 ECN 的瓶颈；路径上有一个“盲区”，闭环就缺一段。",
        "不能修复 CNP 丢失/被 pause/映射错误；反馈断了，RP 会误判。",
        "不能保证 RTT 公平；长 RTT 流可能吃亏。",
    ])
