# -*- coding: utf-8 -*-
"""Front matter and sections 1-3."""
import os
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docx_helpers import (
    NAVY, BLUE, TEAL, ORANGE, RED, GRAY, WHITE,
    LIGHT_BLUE, LIGHT_GRAY, LIGHT_GREEN, LIGHT_ORANGE, LIGHT_RED,
    BODY_FONT, HEAD_FONT,
    P, BULLETS, NUMBERED, H1, H2, H3, CAPTION, TABLE, CODE, CALLOUT, FIGURE,
    set_run_font, set_paragraph_format, add_rich_text,
    set_cell_background, set_cell_margins, set_cell_width,
)
from common import DOC_TITLE, DOC_SUBTITLE, DOC_VERSION, DOC_DATE


def add_cover(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.1)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)
    sec.different_first_page_header_footer = True

    # Top band
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell = t.cell(0, 0)
    set_cell_width(cell, 16.2)
    set_cell_background(cell, NAVY)
    set_cell_margins(cell, 0.28, 0.28, 0.30, 0.30)
    p = cell.paragraphs[0]
    set_paragraph_format(p, before=0, after=0, line=1.0)
    r = p.add_run("NVIDIA NETWORKING   ·   TECHNICAL GUIDE")
    set_run_font(r, HEAD_FONT, 9.5, True, "D9E7F5")

    # Title block
    p = doc.add_paragraph()
    set_paragraph_format(p, before=20, after=4, line=1.12)
    r = p.add_run("ConnectX-6 网卡")
    set_run_font(r, HEAD_FONT, 25, True, NAVY)

    p = doc.add_paragraph()
    set_paragraph_format(p, before=0, after=4, line=1.15)
    r = p.add_run("RoCEv2 拥塞控制与 PFC 流控")
    set_run_font(r, HEAD_FONT, 25, True, NAVY)

    p = doc.add_paragraph()
    set_paragraph_format(p, before=0, after=12, line=1.15)
    r = p.add_run("实施指南")
    set_run_font(r, HEAD_FONT, 25, True, NAVY)

    p = doc.add_paragraph()
    set_paragraph_format(p, before=0, after=4, line=1.3)
    r = p.add_run(DOC_SUBTITLE)
    set_run_font(r, BODY_FONT, 12.0, False, GRAY)

    # thin rule
    p = doc.add_paragraph()
    set_paragraph_format(p, before=4, after=14, line=1.0)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), TEAL)
    pBdr.append(bottom)
    pPr.append(pBdr)

    # Badge table
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    c = t.cell(0, 0)
    set_cell_width(c, 15.8)
    set_cell_background(c, LIGHT_BLUE)
    set_cell_margins(c, 0.18, 0.18, 0.22, 0.22)
    p = c.paragraphs[0]
    set_paragraph_format(p, before=0, after=0, line=1.2)
    r = p.add_run("17 份 NVIDIA 官方文档  ·  DCQCN / ECN / PFC  ·  含命令、验收与排障")
    set_run_font(r, BODY_FONT, 9.6, True, BLUE)

    # Scope
    p = doc.add_paragraph()
    set_paragraph_format(p, before=22, after=6, line=1.2)
    r = p.add_run("适用对象")
    set_run_font(r, HEAD_FONT, 10.0, True, TEAL)
    p = doc.add_paragraph()
    set_paragraph_format(p, before=0, after=0, line=1.45)
    r = p.add_run("网络架构师、RDMA / RoCE 运维工程师、AI/HPC 集群交付与测试人员")
    set_run_font(r, BODY_FONT, 10.2, False, NAVY)

    p = doc.add_paragraph()
    set_paragraph_format(p, before=14, after=6, line=1.2)
    r = p.add_run("适用型号")
    set_run_font(r, HEAD_FONT, 10.0, True, TEAL)
    p = doc.add_paragraph()
    set_paragraph_format(p, before=0, after=0, line=1.45)
    r = p.add_run("ConnectX-6 EN / ConnectX-6 VPI 的 Ethernet 模式；ConnectX-6 Dx/Lx 请按对应固件文档复核参数")
    set_run_font(r, BODY_FONT, 9.8, False, NAVY)

    # Bottom meta
    doc.add_paragraph().paragraph_format.space_after = Pt(28)
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    vals = [("版本", DOC_VERSION), ("日期", DOC_DATE), ("文档状态", "技术整理稿 · 实施前按环境复核")]
    for i, (k, v) in enumerate(vals):
        c = t.cell(0, i)
        set_cell_width(c, 5.4)
        set_cell_margins(c, 0.10, 0.10, 0.12, 0.12)
        p = c.paragraphs[0]
        set_paragraph_format(p, before=0, after=0, line=1.2, align=WD_ALIGN_PARAGRAPH.LEFT)
        r = p.add_run(k + "\n")
        set_run_font(r, BODY_FONT, 8.2, False, GRAY)
        r = p.add_run(v)
        set_run_font(r, BODY_FONT, 9.0, True, NAVY)

    p = doc.add_paragraph()
    set_paragraph_format(p, before=26, after=0, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("本文件为公开 NVIDIA 文档的工程整理，不替代 NVIDIA 官方技术支持或现场验证。")
    set_run_font(r, BODY_FONT, 8.0, False, GRAY)


def add_front_matter(doc):
    doc.add_page_break()
    H2(doc, "文档信息")
    rows = [
        ("文档名称", DOC_TITLE),
        ("副标题", DOC_SUBTITLE),
        ("版本 / 日期", f"{DOC_VERSION} / {DOC_DATE}"),
        ("证据基线", "17 份 NVIDIA 官方文档（详见附录 D），访问日期 2026-10-25"),
        ("适用范围", "ConnectX-6 Ethernet / RoCEv2；Linux 主机；NVIDIA Onyx 或 Cumulus Linux 交换机示例"),
        ("不适用范围", "InfiniBand 模式的 credit/IB CC；Windows/WinOF-2 的注册表配置；非 NVIDIA 交换机的厂商私有实现"),
        ("维护建议", "每季度或固件/驱动升级后复核一次；参数与默认值以设备 query 输出为准"),
    ]
    TABLE(doc, ["项目", "内容"], rows, [3.2, 13.0], font_size=8.7, first_col_bold=True)

    H2(doc, "阅读说明与证据等级")
    P(doc, "本文把“官方事实”“官方默认/示例”“工程建议”分开写。正文中的 [S1]–[S17] 对应附录 D 的 NVIDIA 官方文档；没有引用的操作性建议，均明确标注为“工程建议”，代表基于机制的落地方法，不是 NVIDIA 强制规范。")
    rows = [
        ("官方事实", "可直接在 NVIDIA 文档中找到原文的能力、协议行为、限制条件。", "文中以 [S#] 标注"),
        ("官方默认/示例", "NVIDIA 产品文档给出的默认值、示例配置或参考拓扑；会随版本/平台变化。", "文中写“默认/示例”，并提示复核"),
        ("工程建议", "为实现可运维、可验证的部署而给出的做法；不是协议强制。", "文中写“建议/应/不要”，并解释依据"),
        ("工程公式", "用于量级校验的推导公式，不能替代厂商缓冲计算。", "文中明确标注“工程校验”"),
    ]
    TABLE(doc, ["等级", "含义", "在本文中的标记"], rows, [3.0, 9.6, 3.6], font_size=8.5, first_col_bold=True)

    H2(doc, "目录")
    toc = [
        ("1", "执行摘要", "9", "验证、验收与性能测试"),
        ("2", "范围、术语与证据方法", "10", "运维监控、故障排查与应急"),
        ("3", "CX6 的能力边界与机制定位", "11", "风险、限制与回滚"),
        ("4", "RoCEv2 拥塞控制：ECN 与 DCQCN", "12", "结论与检查清单"),
        ("5", "DCQCN 参数、默认值与调优", "A", "附录 A  命令速查"),
        ("6", "PFC 流控：原理、缓冲与阈值", "B", "附录 B  DCQCN 参数总表"),
        ("7", "端到端设计：优先级、PFC 域与 ECN 域", "C", "附录 C  官方原文证据摘录"),
        ("8", "实施步骤：从基线到上线", "D", "附录 D  参考文档与链接"),
    ]
    TABLE(doc, ["章", "内容", "章", "内容"], toc, [1.2, 6.8, 1.2, 6.8], font_size=8.2, first_col_bold=True, align_center=False)

    CALLOUT(doc, "如何使用本文", "第一次部署：按第 3→7 章做设计，再按第 8 章实施，最后按第 9 章验收。已经在线运行的集群：直接看第 5、6、10 章的调优与排障。需要给评审/审计看依据：看第 2 章的证据方法、附录 C 的原文摘录和附录 D 的链接。", "info")


def add_section_1(doc):
    h = H1(doc, "执行摘要", "1")
    h.paragraph_format.page_break_before = False
    P(doc, "这份文档回答三个问题：ConnectX-6 上的 RoCEv2 拥塞控制与 PFC 分别是什么、应该如何端到端配置、上线后如何证明它真的生效。结论建立在 NVIDIA 公开文档之上；凡是超出文档原文的工程做法，文中都会明确标注。")

    H2(doc, "结论先行", "1.1")
    rows = [
        ("PFC 与 DCQCN 不是替代关系", "PFC 是逐跳、逐优先级的链路反压；DCQCN 是端到端、基于 ECN/CNP 的速率控制。前者防瞬时溢出，后者防持续拥塞。", "[S3][S5][S8]"),
        ("RoCEv2 有两条合法路线", "Lossless：PFC + ECN/DCQCN；Lossy：只靠 ECN/DCQCN + 传输层重传。选哪条由业务对丢包/时延的容忍度决定。", "[S10][S13]"),
        ("NVIDIA 常见默认值不是协议强制值", "RoCE 常用 priority 3 / DSCP 26，CNP 常用 DSCP 48 / PCP 6；Onyx 默认 ECN 150–1500，Cumulus 默认 150 KB–1.5 MB。数值必须端到端一致，不能只看单台设备。", "[S6][S10][S12][S13]"),
        ("PFC 的成败在缓冲与水线", "Xoff 之前要有足够 headroom 吸收在途报文；cable length、MTU、速率、Xoff/Xon 任一错配，都会表现为“配了 PFC 仍然丢包”或“PFC 频繁触发”。", "[S9][S12]"),
        ("CNP 必须被优先转发", "接收端网卡生成 CNP 回送源端；CNP 被丢、被 pause、或被排到数据后面，DCQCN 就会失效或过冲。", "[S5][S6][S10]"),
        ("持久化要单独设计", "mlxconfig 的固件参数可跨重启保留；PFC/trust/ECN sysfs 等运行时配置通常需要在每次启动后重新应用。", "[S14][S15]"),
    ]
    TABLE(doc, ["结论", "要点", "依据"], rows, [4.2, 9.6, 2.4], font_size=8.4, first_col_bold=True)

    H2(doc, "推荐实施路线", "1.2")
    NUMBERED(doc, [
        "**定模式**：业务能否容忍重传/乱序？不能 → lossless（PFC+ECN）；能 → lossy（ECN）；只在边缘吸收微突发 → semi-lossless。",
        "**定优先级**：RoCE 数据、CNP、普通流量分开；把 RoCE 放到独立 lossless priority，CNP 放 strict priority。",
        "**定基线**：记录固件/驱动、`mlxconfig query`、`mlnx_qos`、交换机 `show roce` 和关键计数器。",
        "**先固件、再交换机、后主机**：固件参数（CC 优先级、CNP DSCP/PCP）需要重启；交换机与主机运行时配置在维护窗口内落地。",
        "**用 incast 验收**：单流只能证明“能跑”，多打一才能证明 PFC/ECN 真正工作。",
        "**把配置做成启动项**：PFC、trust、ECN、sysfs 参数纳入 systemd/Ansible/Network Operator，重启后自动复核。",
    ])

    H2(doc, "三条红线", "1.3")
    CALLOUT(doc, "不要做的事", "① 不要只在一端开 PFC；② 不要把 RoCE 和普通 TCP 放在同一个 lossless priority；③ 不要在未确认 CNP 优先级和 ECN 阈值前，靠调大 PFC Xoff 硬扛拥塞。前两条会造成“看似 lossless、实际丢包/风暴”，第三条会把拥塞变成大面积反压。", "risk")
    BULLETS(doc, [
        "不要用 PFC 代替带宽规划：PFC 只能把压力反压回上游，不能创造带宽。",
        "不要把 Onyx 的 150/1500 直接当成字节数，也不要把 Cumulus 的 150000/1500000 直接套到所有平台；先确认单位和 TC。",
        "不要在生产窗口直接改 `LINK_TYPE`、RoCE 模式或 `mlxconfig` 后不重启；这些操作可能中断业务。",
        "不要忽略 `tx_pause_storm_*`、`no-buffer-discard`、`out_of_sequence` 这些“已经在报警”的信号。",
    ])

    H2(doc, "适用读者与阅读路径", "1.4")
    rows = [
        ("网络架构师", "第 3、6、7 章", "确定 lossless/lossy、优先级、PFC 域与缓冲策略"),
        ("交付/实施工程师", "第 8、9 章 + 附录 A", "按步骤配置、验收、留档"),
        ("运维/SRE", "第 5、10、11 章", "监控、排障、变更与回滚"),
        ("评审/审计", "第 2 章 + 附录 C/D", "核对每条结论的官方依据"),
    ]
    TABLE(doc, ["角色", "优先阅读", "关注点"], rows, [4.2, 4.2, 7.8], font_size=8.5, first_col_bold=True)


def add_section_2(doc):
    H1(doc, "范围、术语与证据方法", "2")
    P(doc, "本文讨论的是 Ethernet 模式下的 RoCEv2。ConnectX-6 也支持 InfiniBand；IB 使用基于信用的链路流控和 IB 拥塞控制，不使用 Ethernet PFC/ECN。把本文的 PFC 命令套到 IB 端口上，不会得到预期的 lossless 行为。")

    H2(doc, "适用范围", "2.1")
    rows = [
        ("适用", "ConnectX-6 EN / VPI 的 Ethernet 模式；RoCEv1/v2 的 Ethernet 承载；Linux OFED/DOCA-OFED；NVIDIA Onyx / Cumulus Linux 的 RoCE 配置示例。"),
        ("可参考", "ConnectX-6 Dx / Lx 的同类机制；但本文的型号能力与默认值证据主要来自 ConnectX-6 文档，使用前必须用对应型号的固件文档复核。"),
        ("不适用", "InfiniBand credit/IB CC；Windows/WinOF-2 注册表；非 NVIDIA 交换机的私有 PFC/ECN 实现；应用层拥塞控制（如 TCP BBR）。"),
        ("必须复核", "固件版本、驱动版本、交换机 NOS 版本、端口速率、MTU、线缆长度、是否启用 DCBX、是否 SR-IOV。"),
    ]
    TABLE(doc, ["边界", "说明"], rows, [3.0, 13.2], font_size=8.6, first_col_bold=True)

    H2(doc, "术语表", "2.2")
    rows = [
        ("PFC", "Priority Flow Control，IEEE 802.1Qbb；按 802.1p 优先级独立暂停/恢复的链路流控。", "[S3][S8]"),
        ("ECN", "Explicit Congestion Notification，RFC 3168；用 IP 头 2 bit 标记拥塞，不丢包。", "[S4][S10]"),
        ("CNP", "Congestion Notification Packet；接收端看到 CE 标记后回送源端的拥塞通知报文。", "[S5]"),
        ("DCQCN", "Data Center Quantized Congestion Notification；NVIDIA 网卡上的 ECN 闭环速率控制算法。", "[S5][S7]"),
        ("CP", "Congestion Point，通常是以太网交换机 Egress 队列。", "[S5]"),
        ("NP", "Notification Point，接收端网卡；生成 CNP。", "[S5]"),
        ("RP", "Reaction Point，发送端网卡；按 CNP 调整 QP 速率。", "[S5]"),
        ("Kmin / Kmax / Pmax", "ECN 标记的起始阈值、全标记阈值、最大标记概率。", "[S5]"),
        ("Xoff / Xon", "PFC 的暂停/恢复水线；Xoff 触发 Pause，Xon 停止 Pause。", "[S8][S9]"),
        ("Headroom", "为 Xoff 之后仍在链路上的报文预留的入口缓冲。", "[S9][S12]"),
        ("PG / TC / SP", "Priority Group / Traffic Class / Switch Priority；分别对应入口缓冲组、出口队列、交换机内部优先级。", "[S9]"),
        ("DSCP / PCP", "IP 头优先级 / VLAN 标签优先级；L3 与 L2 两种标记方式。", "[S2][S17]"),
        ("DCBX", "Data Center Bridging Exchange；用 LLDP 把 PFC/ETS 配置从交换机下发到主机。", "[S2][S3]"),
        ("BDP", "Bandwidth-Delay Product，带宽时延积；headroom 估算的核心量。", "工程定义"),
        ("HOL", "Head-of-Line blocking，队头阻塞；同一优先级内的一条慢流会挡住后续报文。", "工程定义"),
        ("PFC storm", "暂停帧沿路径扩散、缓冲区连锁占满，导致大面积停传。", "[S8]"),
    ]
    TABLE(doc, ["术语", "含义", "依据"], rows, [3.0, 10.6, 2.6], font_size=8.2, first_col_bold=True)

    H2(doc, "证据方法", "2.3")
    P(doc, "本文只把 NVIDIA 官方文档作为“事实”来源，优先使用产品手册、MLNX_OFED/DOCA 文档、Onyx/Cumulus 用户手册和 MFT 文档。厂商博客、第三方文章只在明显标注时作为背景，不作为参数默认值或命令语义的依据。")
    rows = [
        ("S1–S4", "网卡能力、主机 QoS/PFC/ECN 配置", "产品手册 + OFED 用户手册"),
        ("S5–S7", "DCQCN 算法、参数、CC 基础设施", "DOCA 文档"),
        ("S8–S11, S17", "交换机 PFC/缓冲/RoCE 行为与默认值", "Onyx 用户手册"),
        ("S12–S13", "交换机 RoCE 默认配置、PFC watchdog、ECN", "Cumulus Linux 用户手册"),
        ("S14–S16", "固件配置、Network Operator 自动化、计数器路径", "MFT / Network Operator / DOCA 文档"),
    ]
    TABLE(doc, ["来源组", "用于本文的内容", "文档类型"], rows, [3.0, 9.2, 4.0], font_size=8.4, first_col_bold=True)
    P(doc, "所有链接在 2026-10-25 访问。NVIDIA 文档可能要求登录企业支持账号；本文引用的页面均为公开可访问页面。")

    H2(doc, "版本基线", "2.4")
    rows = [
        ("网卡", "ConnectX-6 EN/VPI（Ethernet 模式）", "以 `mlxconfig query`、`ethtool -i`、PSID 为准"),
        ("主机软件", "MLNX_OFED 24.10-2.1.8 LTS 或 DOCA-OFED 对应版本", "命令语义以本机 `--help` 为准"),
        ("固件工具", "MFT 4.36.0", "mlxconfig 需要 mst driver"),
        ("交换机示例", "Onyx 3.10.4800；Cumulus Linux 5.15", "不同 NOS 命令不同，不要混用"),
        ("文档日期", "2026-10-25", "升级后重新核对默认值"),
    ]
    TABLE(doc, ["层", "本文基线", "复核方法"], rows, [3.0, 8.6, 4.6], font_size=8.5, first_col_bold=True)

    H2(doc, "三个容易混淆的概念", "2.5")
    rows = [
        ("QCN（802.1Qau）", "Layer 2 拥塞通知，使用 CNM；CX6 手册列有该能力。", "不是 RoCEv2 的 DCQCN；不要用 QCN 的配置思路替代 ECN/CNP。", "[S1][S5]"),
        ("DCQCN", "Layer 3 ECN + CNP 的端到端速率控制；NVIDIA 网卡硬件实现。", "RoCEv2 的主流 CC；依赖交换机 ECN 标记和 CNP 优先转发。", "[S5][S7]"),
        ("PFC", "Layer 2 逐跳、逐优先级的暂停/恢复。", "不是拥塞控制；不降低发送速率，只把压力反压回上游。", "[S8]"),
        ("IB CC", "InfiniBand 的 credit/拥塞控制。", "与 Ethernet PFC/ECN 不同栈；VPI 端口切到 IB 后不要套用本文。", "工程边界"),
    ]
    TABLE(doc, ["概念", "是什么", "与本文的关系", "依据"], rows, [3.2, 4.6, 5.8, 2.6], font_size=8.2, first_col_bold=True)


def add_section_3(doc):
    H1(doc, "CX6 的能力边界与机制定位", "3")
    P(doc, "ConnectX-6 是双端口 200Gb/s 级别的 VPI/以太网卡。就本文主题而言，关键不是峰值带宽，而是它同时具备 PFC、ECN、DCQCN 和 QoS 映射能力，并且这些能力分别落在不同的硬件/协议层。")

    H2(doc, "CX6 与本文相关的能力", "3.1")
    rows = [
        ("以太网速率", "1/10/25/40/50/100/200 Gb/s；双端口 QSFP56（具体型号而异）", "[S1]"),
        ("PFC", "IEEE 802.1Qbb，按优先级暂停/恢复", "[S1]"),
        ("QCN", "IEEE 802.1Qau，Layer 2 拥塞通知（注意不是 DCQCN）", "[S1]"),
        ("ETS / QoS", "IEEE 802.1Qaz；端口级 QoS、优先级映射", "[S1]"),
        ("RoCE", "IBTA RDMA over Converged Ethernet；支持 L2/L3；利用 DCB 与高级拥塞控制硬件", "[S1]"),
        ("ECN / DCQCN", "RoCEv2 的 ECN 标记、CNP 生成/处理、RP 速率调整由网卡硬件参与", "[S4][S5]"),
        ("VPI", "同一卡可配置 IB 或 Ethernet；端口模式由 `LINK_TYPE_Px` 决定", "[S1]"),
        ("管理", "BMC、MCTP、NC-SI 等（与本文无直接关系）", "[S1]"),
    ]
    TABLE(doc, ["能力", "说明", "依据"], rows, [3.2, 10.8, 2.2], font_size=8.5, first_col_bold=True)

    CALLOUT(doc, "型号边界", "本文的默认值证据来自 ConnectX-6 EN/VPI 文档。ConnectX-6 Dx/Lx 的 PFC/ECN 机制同属 Ethernet 范畴，但固件参数名、默认值和特性集合可能不同；不要因为“都是 CX6”就直接复制 `mlxconfig` 值。先用 `mlxconfig query` 和对应型号的 release notes 核对。", "warn")

    H2(doc, "一张图看清三个平面", "3.2")
    FIGURE(doc, os.path.join("assets", "fig_architecture.png"), 16.0)
    CAPTION(doc, "图 1  RoCEv2 中 DCQCN 三角色与 PFC 逐跳反压的关系（根据 [S5][S8][S10] 绘制，概念图非按比例）")

    P(doc, "这张图要表达的不是“PFC 和 DCQCN 一起工作”这么简单，而是三个不同的平面：")
    rows = [
        ("数据平面", "RoCEv2 数据报文", "源端 QP → 交换机 Egress → 目的端 QP", "受 PFC 暂停、受 ECN 标记"),
        ("控制平面", "CNP 反馈", "目的端 NP → 源端 RP", "必须低时延、strict priority、不能被 PFC 误伤"),
        ("流控平面", "PFC Pause/Xon", "每一跳的接收侧 → 发送侧", "逐跳、逐优先级、不跨路由域提供端到端语义"),
    ]
    TABLE(doc, ["平面", "载体", "方向", "关键约束"], rows, [3.0, 4.4, 4.6, 4.0], font_size=8.4, first_col_bold=True)

    H2(doc, "PFC 不是拥塞控制", "3.3")
    P(doc, "这是全文最重要的一句话。PFC 做的是“我这家门口堵了，你先别发”；它不告诉发送端“应该降到多少”，也不感知整条路径的拥塞。DCQCN 做的是“路径上有拥塞，按 QP 把速率降下来”，它通过 ECN 标记和 CNP 形成闭环。两者解决的是不同时间尺度的问题。")
    rows = [
        ("作用范围", "逐跳链路（L2）", "端到端路径（L3）"),
        ("触发信号", "接收侧入口缓冲超过 Xoff", "交换机 Egress 队列超过 Kmin，标记 CE"),
        ("反馈方式", "Pause 帧 / Xon 帧", "CNP 回送源端"),
        ("执行者", "链路两端 MAC/交换芯片", "发送端网卡 RP 的限速器"),
        ("时间尺度", "微秒级、瞬时", "数十微秒到毫秒级、持续"),
        ("主要风险", "反压扩散、PFC storm、HOL", "阈值/参数不当导致降速过度或不足"),
    ]
    TABLE(doc, ["维度", "PFC", "DCQCN / ECN"], rows, [3.4, 6.2, 6.6], font_size=8.4, first_col_bold=True)

    CALLOUT(doc, "Lossless 不等于零丢包", "NVIDIA Onyx 文档明确写道：lossless 并不意味着报文绝对不会丢失。PFC 只能降低入口缓冲溢出的概率；如果 headroom 不足、线缆长度配置错误、PFC 被绕过或发生风暴，仍然会丢包。不要把“开了 PFC”当成丢包问题的终点。", "risk")

    H2(doc, "RoCEv1、RoCEv2 与 IB 的区别", "3.4")
    BULLETS(doc, [
        "**RoCEv2**：IB 传输层跑在 UDP/IP 上（UDP 4791），可路由；本文的 ECN/DCQCN 讨论以 RoCEv2 为主。[S13]",
        "**RoCEv1**：IB 传输层直接跑在 Ethernet 上，Layer 2、不可路由；PFC 仍可用于保证 lossless，但 DCQCN 的 ECN/CNP 机制按 RoCEv2 实现来讨论。[S4][S10]",
        "**IB**：使用基于信用的链路流控和 IB 拥塞控制；没有 Ethernet PFC，也不使用本文的 ECN/DCQCN 参数。VPI 卡切到 IB 后，必须按 IB 文档配置。",
        "**QCN**：CX6 支持 802.1Qau，但它使用 Layer 2 CNM，不是 RoCEv2 的 CNP；不要把 QCN 的调参当成 DCQCN 调参。[S1][S5]",
    ])
